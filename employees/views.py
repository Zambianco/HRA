from calendar import monthrange
from urllib.parse import urlencode
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import date, datetime, time, timedelta
import csv
import ast
import io
import json
import os
import re
import subprocess
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from django.contrib import messages
from django.db import IntegrityError, transaction
from django.db.utils import OperationalError, ProgrammingError
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from .models import (
    Area,
    Employee,
    EmployeeEvent,
    BankHoursRule,
    EmployeeTimeEntry,
    BankHoursSemesterClosure,
    BankHoursSemesterClosureAdjustment,
    TimesheetImportReport,
    TimesheetImportReportRow,
    TimesheetMonthClosure,
    TimesheetMonthClosureCalendarPeriodSnapshot,
    TimesheetMonthClosureEmployeeEventSnapshot,
    TimesheetMonthClosureWorkScheduleSnapshot,
    Cargo,
    Department,
    Sector,
    WorkCalendar,
    WorkCalendarPeriod,
    WorkSchedule,
)

WEEKDAY_HOUR_FIELDS = (
    ("horas_segunda", "segunda-feira"),
    ("horas_terca", "terca-feira"),
    ("horas_quarta", "quarta-feira"),
    ("horas_quinta", "quinta-feira"),
    ("horas_sexta", "sexta-feira"),
    ("horas_sabado", "sabado"),
    ("horas_domingo", "domingo"),
)
SCHEDULE_FIELDS_IN_WEEKDAY_ORDER = tuple(
    field_name for field_name, _ in WEEKDAY_HOUR_FIELDS
)

TIME_ENTRY_MINUTE_FIELDS = (
    ("regular_minutes", "horas normais"),
    ("overtime_60_minutes", "horas extras 60%"),
    ("overtime_100_minutes", "horas extras 100%"),
    ("absence_unexcused_minutes", "ausencias por falta"),
    ("absence_excused_minutes", "ausencias justificadas"),
    ("absence_bank_minutes", "ausencias por banco"),
)


def _project_root_path() -> Path:
    return Path(__file__).resolve().parent.parent


def _default_sqlite_db_path() -> Path:
    return _project_root_path() / "data" / "rh.db"


def _normalize_db_file_path(path_value: str) -> str:
    raw_value = str(path_value or "").strip()
    if not raw_value:
        return str(_default_sqlite_db_path())

    project_root = _project_root_path()
    candidate = Path(raw_value).expanduser()
    if not candidate.is_absolute():
        candidate = project_root / candidate

    try:
        resolved = candidate.resolve(strict=False)
        _ = resolved.parent.exists()
    except OSError:
        return str(_default_sqlite_db_path())

    return str(resolved)


def _portable_db_path_for_storage(path_value: str) -> str:
    normalized = Path(_normalize_db_file_path(path_value))
    project_root = _project_root_path()
    try:
        relative = normalized.relative_to(project_root)
    except ValueError:
        return str(normalized)
    return relative.as_posix()


def _db_config_file_path() -> Path:
    return _project_root_path() / "data" / "db_config.json"


def _load_db_runtime_config() -> dict:
    config_path = _db_config_file_path()
    if not config_path.exists():
        return {}
    try:
        loaded = json.loads(config_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    if not isinstance(loaded, dict):
        return {}

    mode = str(loaded.get("mode", "")).strip().lower()
    if mode not in {"online", "arquivo"}:
        mode = "arquivo"

    normalized = {
        "mode": mode,
        "db_file_path": _normalize_db_file_path(str(loaded.get("db_file_path", ""))),
    }
    return normalized


def _save_db_runtime_config(mode: str, db_file_path: str) -> None:
    config_path = _db_config_file_path()
    config_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "mode": mode,
        "db_file_path": _portable_db_path_for_storage(db_file_path),
    }
    config_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _open_location_in_file_manager(path_value: str) -> None:
    target = Path(path_value).expanduser()
    location = target.parent if target.suffix.lower() == ".db" else target
    resolved = location.resolve()

    if os.name == "nt":
        os.startfile(str(resolved))
        return
    if os.name == "posix":
        subprocess.Popen(["xdg-open", str(resolved)])
        return
    raise OSError("Sistema operacional nao suportado para abrir localizacao.")


def _create_clean_sqlite_db(db_file_path: str) -> tuple[bool, str]:
    target = Path(_normalize_db_file_path(db_file_path))
    parent = target.parent
    if not parent.exists():
        parent.mkdir(parents=True, exist_ok=True)

    project_root = _project_root_path()
    env = os.environ.copy()
    env["DATABASE_MODE"] = "arquivo"
    env["DB_FILE_PATH"] = str(target)

    migrate_cmd = [sys.executable, "manage.py", "migrate", "--noinput"]
    completed = subprocess.run(
        migrate_cmd,
        cwd=str(project_root),
        env=env,
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0:
        stderr = (completed.stderr or "").strip()
        stdout = (completed.stdout or "").strip()
        detail = stderr or stdout or "Falha desconhecida durante migracao."
        return False, detail[:800]
    return True, ""


def _employees_redirect_with_flags(**params):
    base_url = reverse("employees_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _events_redirect_with_filters(employee_id=None, event_type=None):
    params = {}
    if employee_id:
        params["employee_id"] = employee_id
    if event_type:
        params["event_type"] = event_type

    base_url = reverse("events_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _timesheet_redirect_with_filters(competence_month=None):
    params = {}
    if competence_month:
        params["competence_month"] = competence_month

    base_url = reverse("timesheet_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _bank_hours_redirect_with_filters(start_month=None, end_month=None):
    params = {}
    if start_month:
        params["start_month"] = start_month
    if end_month:
        params["end_month"] = end_month

    base_url = reverse("bank_hours_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _bank_hours_rules_redirect_with_filters(start_month=None, end_month=None):
    params = {}
    if start_month:
        params["start_month"] = start_month
    if end_month:
        params["end_month"] = end_month

    base_url = reverse("bank_hours_rules_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _snapshots_audit_redirect_with_filters(competence_month=None):
    params = {}
    if competence_month:
        params["competence_month"] = competence_month

    base_url = reverse("timesheet_snapshots_audit_page")
    query = urlencode(params)
    return f"{base_url}?{query}" if query else base_url


def _employee_edit_redirect_with_tab(employee_id, tab=None):
    base_url = reverse("employee_edit_page", kwargs={"employee_id": employee_id})
    query = urlencode({"tab": tab}) if tab else ""
    return f"{base_url}?{query}" if query else base_url


def _parse_daily_hours(request_data):
    parsed_hours = {}

    for field_name, label in WEEKDAY_HOUR_FIELDS:
        raw_value = (request_data.get(field_name) or "").strip().replace(",", ".")

        if not raw_value:
            return None, f'Informe a carga horaria para "{label}".'

        try:
            hours = Decimal(raw_value)
        except InvalidOperation:
            return None, f'Valor invalido de carga horaria para "{label}".'

        if hours < 0 or hours > 24:
            return None, f'A carga horaria de "{label}" deve estar entre 0 e 24.'

        parsed_hours[field_name] = hours

    return parsed_hours, None


def _parse_special_period_dates(request_data):
    start_date_raw = (request_data.get("start_date") or "").strip()
    end_date_raw = (request_data.get("end_date") or "").strip()

    if not start_date_raw:
        return None, None, 'Informe a data de inicio do periodo.'

    if not end_date_raw:
        return None, None, 'Informe a data de fim do periodo.'

    try:
        start_date = date.fromisoformat(start_date_raw)
    except ValueError:
        return None, None, 'Informe uma data de inicio valida.'

    try:
        end_date = date.fromisoformat(end_date_raw)
    except ValueError:
        return None, None, 'Informe uma data de fim valida.'

    if end_date < start_date:
        return None, None, 'A data de fim deve ser maior ou igual a data de inicio.'

    return start_date, end_date, None


def _parse_non_negative_minutes(raw_value, field_label):
    normalized_value = (raw_value or "").strip()
    if not normalized_value:
        return 0, None

    compact_value = normalized_value.replace(" ", "")
    value = None

    if compact_value.isdigit():
        value = int(compact_value)
    else:
        hhmm_match = re.fullmatch(r"(\d+):([0-5]?\d)", compact_value)
        if not hhmm_match:
            return None, f'Valor invalido para "{field_label}".'
        hours = int(hhmm_match.group(1))
        minutes = int(hhmm_match.group(2))
        value = (hours * 60) + minutes

    if value < 0:
        return None, f'O valor de "{field_label}" nao pode ser negativo.'

    if value > 59999:
        return None, f'O valor de "{field_label}" deve ser no maximo 59999 minutos.'

    return value, None



def _time_value_to_decimal_hours(value):
    if value is None or value == "":
        return 0.0

    if isinstance(value, timedelta):
        return round(value.total_seconds() / 3600, 2)

    if isinstance(value, time):
        return round(value.hour + (value.minute / 60) + (value.second / 3600), 2)

    if isinstance(value, datetime):
        return round(value.hour + (value.minute / 60) + (value.second / 3600), 2)

    raw = str(value).strip()
    if not raw:
        return 0.0

    if "day" in raw:
        day_parts = raw.split(",")
        days = int(day_parts[0].split()[0])
        hour_text = day_parts[1].strip() if len(day_parts) > 1 else "0:00:00"
        h, m, s = map(int, hour_text.split(":"))
        return round((days * 24) + h + (m / 60) + (s / 3600), 2)

    parts = raw.split(":")
    h = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 else 0
    s = int(parts[2]) if len(parts) > 2 else 0
    return round(h + (m / 60) + (s / 3600), 2)


def _decimal_hours_to_minutes(hours_value):
    return int(round(Decimal(str(hours_value)) * Decimal("60")))



def _parse_point_sheet_header_period(raw_text):
    text = (str(raw_text or "")).strip()
    if not text:
        return None, None, "Cabecalho A2 vazio: informe periodo no formato De: DD/MM/AAAA ate DD/MM/AAAA."

    match = re.search(r"de\s*:\s*(\d{2}/\d{2}/\d{4})\s*at[eé]\s*(\d{2}/\d{2}/\d{4})", text, re.IGNORECASE)
    if not match:
        return None, None, "Nao foi possivel ler o periodo em A2. Formato esperado: De: DD/MM/AAAA ate DD/MM/AAAA."

    try:
        start_date = datetime.strptime(match.group(1), "%d/%m/%Y").date()
        end_date = datetime.strptime(match.group(2), "%d/%m/%Y").date()
    except ValueError:
        return None, None, "Periodo em A2 invalido."

    if end_date < start_date:
        return None, None, "Periodo em A2 invalido: data final menor que inicial."

    return start_date, end_date, None
def _build_timesheet_import_preview(uploaded_file, competence_month):
    if not uploaded_file:
        return None, ["Selecione um arquivo XLSX para importar."], []

    source_file_name = uploaded_file.name or ""
    file_name = source_file_name.lower()
    if not file_name.endswith(".xlsx"):
        return None, ["Formato invalido. Envie um arquivo .xlsx."], []

    try:
        workbook = load_workbook(uploaded_file, data_only=True)
    except Exception as exc:
        return None, [f"Nao foi possivel ler o arquivo XLSX: {exc}"], []

    first_sheet = workbook.worksheets[0] if workbook.worksheets else None
    if not first_sheet:
        return None, ["Arquivo XLSX sem abas para importacao."], []

    period_start, period_end, period_error = _parse_point_sheet_header_period(first_sheet["A2"].value)
    if period_error:
        return None, [period_error], []

    expected_start, expected_end = _get_month_date_range(competence_month)
    if period_start != expected_start or period_end != expected_end:
        return None, [
            (
                "Periodo do espelho ponto diferente do mes selecionado. "
                f"Arquivo: {period_start.strftime('%d/%m/%Y')} ate {period_end.strftime('%d/%m/%Y')}. "
                f"Competencia: {expected_start.strftime('%d/%m/%Y')} ate {expected_end.strftime('%d/%m/%Y')}."
            )
        ], []

    aggregated_by_registration = defaultdict(
        lambda: {
            "regular_minutes": 0,
            "overtime_60_minutes": 0,
            "overtime_100_minutes": 0,
            "absence_unexcused_minutes": 0,
            "absence_excused_minutes": 0,
        }
    )
    warnings = []
    hard_errors = []

    for sheet in workbook.worksheets:
        registration = None
        for row in range(1, sheet.max_row + 1):
            value_a = sheet[f"A{row}"].value
            value_a = str(value_a).strip() if value_a else ""

            if value_a == "Nº Folha":
                registration = (sheet[f"B{row}"].value or "")
                registration = str(registration).strip()
                continue

            if value_a != "Data":
                continue

            if not registration:
                warnings.append(
                    f'Aba "{sheet.title}" linha {row}: bloco sem matricula em "Nº Folha".'
                )
                continue

            values_row = row + 1
            try:
                normal_b = _time_value_to_decimal_hours(sheet[f"B{values_row}"].value)
                normal_g = _time_value_to_decimal_hours(sheet[f"G{values_row}"].value)
                absence_unexcused = _time_value_to_decimal_hours(sheet[f"C{values_row}"].value)
                overtime_60_d = _time_value_to_decimal_hours(sheet[f"D{values_row}"].value)
                overtime_60_f = _time_value_to_decimal_hours(sheet[f"F{values_row}"].value)
                overtime_100 = _time_value_to_decimal_hours(sheet[f"E{values_row}"].value)
                absence_excused = _time_value_to_decimal_hours(sheet[f"G{values_row}"].value)
            except Exception as exc:
                warnings.append(
                    f'Aba "{sheet.title}" linha {values_row}: valor invalido ({exc}).'
                )
                continue

            regular_hours = normal_b - normal_g
            overtime_60_hours = overtime_60_d + overtime_60_f

            entry = aggregated_by_registration[registration]
            entry["regular_minutes"] += _decimal_hours_to_minutes(regular_hours)
            entry["overtime_60_minutes"] += _decimal_hours_to_minutes(overtime_60_hours)
            entry["overtime_100_minutes"] += _decimal_hours_to_minutes(overtime_100)
            entry["absence_unexcused_minutes"] += _decimal_hours_to_minutes(absence_unexcused)
            entry["absence_excused_minutes"] += _decimal_hours_to_minutes(absence_excused)

    if not aggregated_by_registration:
        return None, ["Nenhum bloco de horas valido foi encontrado no arquivo."], warnings

    base_employees = _get_timesheet_base_employees()
    base_employee_ids = [employee.id for employee in base_employees]
    lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
    eligible_employees = _filter_employees_active_in_period(
        base_employees,
        expected_start,
        expected_end,
        lifecycle_dates,
    )
    eligible_employee_ids = {employee.id for employee in eligible_employees}
    employees_by_registration = {
        str(employee.matricula).strip(): employee
        for employee in base_employees
        if str(employee.matricula or "").strip()
    }
    preview_rows = []
    skipped_rows = []
    imported_registrations = {
        str(registration or "").strip()
        for registration in aggregated_by_registration.keys()
        if str(registration or "").strip()
    }
    for registration, values in aggregated_by_registration.items():
        issues = []
        employee = employees_by_registration.get(registration)
        if not employee:
            issues.append("matricula nao encontrada entre empregados ativos")
        elif employee.id not in eligible_employee_ids:
            issues.append("empregado fora da competencia pela data de admissao/demissao")

        for field_name, field_label in TIME_ENTRY_MINUTE_FIELDS:
            if field_name == "absence_bank_minutes":
                continue
            value = values[field_name]
            _, value_error = _parse_non_negative_minutes(str(value), field_label)
            if value_error:
                issues.append(value_error)

        absence_unexcused_minutes = values["absence_unexcused_minutes"]
        absence_bank_minutes = 0
        if (
            employee
            and employee.regime_compensacao_jornada
            == Employee.REGIME_COMPENSACAO_PARTICIPANTE
        ):
            absence_bank_minutes = absence_unexcused_minutes
            absence_unexcused_minutes = 0

        row_payload = {
            "registration": registration,
            "employee_id": employee.id if employee else None,
            "employee_name": employee.nome_completo if employee else "",
            "regular_minutes": values["regular_minutes"],
            "overtime_60_minutes": values["overtime_60_minutes"],
            "overtime_100_minutes": values["overtime_100_minutes"],
            "absence_unexcused_minutes": absence_unexcused_minutes,
            "absence_excused_minutes": values["absence_excused_minutes"],
            "absence_bank_minutes": absence_bank_minutes,
            "issues": issues,
        }
        if issues:
            skipped_rows.append(row_payload)
        else:
            preview_rows.append(row_payload)

    missing_employee_rows = []
    for employee in eligible_employees:
        registration = str(employee.matricula or "").strip()
        if registration and registration in imported_registrations:
            continue

        issues = []
        if registration:
            issues.append("empregado no timesheet da competencia nao encontrado na importacao")
        else:
            issues.append("empregado no timesheet sem matricula cadastrada")

        missing_employee_rows.append(
            {
                "registration": registration,
                "employee_id": employee.id,
                "employee_name": employee.nome_completo,
                "regular_minutes": 0,
                "overtime_60_minutes": 0,
                "overtime_100_minutes": 0,
                "absence_unexcused_minutes": 0,
                "absence_excused_minutes": 0,
                "absence_bank_minutes": 0,
                "issues": issues,
            }
        )

    if not preview_rows:
        hard_errors.append("Nenhuma linha valida para importacao apos validacao.")

    preview = {
        "source_file_name": source_file_name,
        "period_start": period_start,
        "period_end": period_end,
        "rows": preview_rows,
        "skipped_rows": skipped_rows,
        "missing_employee_rows": missing_employee_rows,
        "warnings": warnings,
        "hard_errors": hard_errors,
    }
    return preview, hard_errors, warnings


def _as_list(value):
    return value if isinstance(value, list) else []


def _create_timesheet_import_report(preview, competence_month):
    rows = preview.get("rows") or []
    skipped_rows = preview.get("skipped_rows") or []
    missing_employee_rows = preview.get("missing_employee_rows") or []
    warnings = _as_list(preview.get("warnings"))
    hard_errors = _as_list(preview.get("hard_errors"))

    report = TimesheetImportReport.objects.create(
        competence_month=competence_month,
        source_file_name=(preview.get("source_file_name") or "")[:255],
        file_period_start=preview.get("period_start"),
        file_period_end=preview.get("period_end"),
        summary={
            "warnings": warnings,
            "hard_errors": hard_errors,
            "valid_rows_count": len(rows),
            "skipped_rows_count": len(skipped_rows),
            "missing_employee_rows_count": len(missing_employee_rows),
        },
    )

    report_rows = []

    def append_report_row(row, row_type):
        issues = row.get("issues") or []
        if not isinstance(issues, list):
            issues = [str(issues)]
        report_rows.append(
            TimesheetImportReportRow(
                report=report,
                row_type=row_type,
                sort_order=len(report_rows) + 1,
                employee_id=row.get("employee_id") or None,
                registration=str(row.get("registration") or "")[:50],
                employee_name=str(row.get("employee_name") or "")[:150],
                regular_minutes=int(row.get("regular_minutes") or 0),
                overtime_60_minutes=int(row.get("overtime_60_minutes") or 0),
                overtime_100_minutes=int(row.get("overtime_100_minutes") or 0),
                absence_unexcused_minutes=int(row.get("absence_unexcused_minutes") or 0),
                absence_excused_minutes=int(row.get("absence_excused_minutes") or 0),
                absence_bank_minutes=int(row.get("absence_bank_minutes") or 0),
                issues=issues,
            )
        )

    for row in rows:
        append_report_row(row, TimesheetImportReportRow.TYPE_IMPORTED)
    for row in skipped_rows:
        append_report_row(row, TimesheetImportReportRow.TYPE_SKIPPED)
    for row in missing_employee_rows:
        append_report_row(row, TimesheetImportReportRow.TYPE_MISSING_IN_IMPORT)

    if report_rows:
        TimesheetImportReportRow.objects.bulk_create(report_rows)

    return report


def _timesheet_import_report_messages(report):
    summary = report.summary if isinstance(report.summary, dict) else {}
    return _as_list(summary.get("warnings")), _as_list(summary.get("hard_errors"))


def _decorate_timesheet_import_report_rows(rows):
    decorated_rows = []
    for row in rows:
        total_minutes = (
            row.regular_minutes
            + row.overtime_60_minutes
            + row.overtime_100_minutes
            + row.absence_unexcused_minutes
            + row.absence_excused_minutes
            + row.absence_bank_minutes
        )
        decorated_rows.append(
            {
                "row": row,
                "issues": _as_list(row.issues),
                "regular_label": _format_minutes_as_hour_label(row.regular_minutes),
                "overtime_60_label": _format_minutes_as_hour_label(row.overtime_60_minutes),
                "overtime_100_label": _format_minutes_as_hour_label(row.overtime_100_minutes),
                "absence_unexcused_label": _format_minutes_as_hour_label(row.absence_unexcused_minutes),
                "absence_excused_label": _format_minutes_as_hour_label(row.absence_excused_minutes),
                "absence_bank_label": _format_minutes_as_hour_label(row.absence_bank_minutes),
                "total_label": _format_minutes_as_hour_label(total_minutes),
            }
        )
    return decorated_rows


def _apply_timesheet_import_report(report):
    expected_start, expected_end = _get_month_date_range(report.competence_month)
    if report.file_period_start != expected_start or report.file_period_end != expected_end:
        return 0, (
            "Periodo do relatorio diferente do mes selecionado. "
            f"Arquivo: {report.file_period_start} ate {report.file_period_end}. "
            f"Competencia: {expected_start} ate {expected_end}."
        )

    _, hard_errors = _timesheet_import_report_messages(report)
    if hard_errors:
        return 0, "Relatorio contem erro bloqueante e nao pode ser aprovado."

    import_rows = list(
        report.rows.select_related("employee").filter(
            row_type=TimesheetImportReportRow.TYPE_IMPORTED
        )
    )
    if not import_rows:
        return 0, "Relatorio sem lancamentos aptos para importar."

    month_start, month_end = _get_month_date_range(report.competence_month)
    base_employees = _get_timesheet_base_employees()
    base_employee_ids = [employee.id for employee in base_employees]
    lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
    employees = _filter_employees_active_in_period(
        base_employees,
        month_start,
        month_end,
        lifecycle_dates,
    )
    employee_by_id = {employee.id: employee for employee in employees}

    is_bank_hours_column_locked = _is_month_locked_by_bank_hours_semester_closure(
        report.competence_month
    )
    entries_in_month = {
        entry.employee_id: entry
        for entry in EmployeeTimeEntry.objects.filter(
            competence_month=report.competence_month
        ).select_related("employee")
    }

    imported_count = 0
    for row in import_rows:
        if not row.employee_id:
            continue
        employee = employee_by_id.get(row.employee_id)
        if not employee:
            continue

        existing_entry = entries_in_month.get(employee.id)
        defaults = {
            "regular_minutes": int(row.regular_minutes or 0),
            "overtime_60_minutes": int(row.overtime_60_minutes or 0),
            "overtime_100_minutes": int(row.overtime_100_minutes or 0),
            "absence_unexcused_minutes": int(row.absence_unexcused_minutes or 0),
            "absence_excused_minutes": int(row.absence_excused_minutes or 0),
            "absence_bank_minutes": int(row.absence_bank_minutes or 0),
            "notes": existing_entry.notes if existing_entry else "",
        }
        if (
            employee.regime_compensacao_jornada
            != Employee.REGIME_COMPENSACAO_PARTICIPANTE
        ):
            defaults["absence_bank_minutes"] = 0
        if (
            is_bank_hours_column_locked
            and employee.regime_compensacao_jornada == Employee.REGIME_COMPENSACAO_PARTICIPANTE
        ):
            defaults["absence_bank_minutes"] = (
                existing_entry.absence_bank_minutes if existing_entry else 0
            )

        has_content = any(
            defaults[field_name] > 0
            for field_name, _ in TIME_ENTRY_MINUTE_FIELDS
        ) or defaults["notes"]
        if not has_content:
            continue

        EmployeeTimeEntry.objects.update_or_create(
            employee=employee,
            competence_month=report.competence_month,
            defaults=defaults,
        )
        imported_count += 1

    if not imported_count:
        return 0, "Nenhum lancamento do relatorio foi aplicado ao timesheet."

    return imported_count, None


def _resolve_competence_month(raw_value):
    normalized = (raw_value or "").strip()
    if not normalized:
        today = timezone.localdate()
        month_start = date(today.year, today.month, 1)
        return month_start, month_start.strftime("%Y-%m"), None

    try:
        month_start = date.fromisoformat(f"{normalized}-01")
    except ValueError:
        return None, None, "Informe um mes de competencia valido."

    return month_start, normalized, None


def _resolve_competence_month_interval(raw_start, raw_end, raw_fallback=None):
    normalized_start = (raw_start or "").strip()
    normalized_end = (raw_end or "").strip()
    normalized_fallback = (raw_fallback or "").strip()

    # Backward compatibility for existing links that still pass competence_month.
    if not normalized_start and not normalized_end:
        fallback_month, fallback_value, error = _resolve_competence_month(normalized_fallback)
        if error:
            return None, None, None, None, error
        return fallback_month, fallback_month, fallback_value, fallback_value, None

    if not normalized_start:
        normalized_start = normalized_end
    if not normalized_end:
        normalized_end = normalized_start

    start_month, start_value, start_error = _resolve_competence_month(normalized_start)
    if start_error:
        return None, None, None, None, "Informe um mes inicial valido."

    end_month, end_value, end_error = _resolve_competence_month(normalized_end)
    if end_error:
        return None, None, None, None, "Informe um mes final valido."

    if end_month < start_month:
        return None, None, None, None, "O mes final deve ser maior ou igual ao mes inicial."

    return start_month, end_month, start_value, end_value, None


def _safe_decimal_formula_eval(expression, variables):
    normalized = (expression or "").strip()
    if not normalized:
        raise ValueError("Formula vazia.")
    normalized = normalized.replace(",", ".")
    parsed = ast.parse(normalized, mode="eval")

    def _eval(node):
        if isinstance(node, ast.Expression):
            return _eval(node.body)
        if isinstance(node, ast.BinOp):
            left = _eval(node.left)
            right = _eval(node.right)
            if isinstance(node.op, ast.Add):
                return left + right
            if isinstance(node.op, ast.Sub):
                return left - right
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Div):
                if right == 0:
                    raise ValueError("Divisao por zero.")
                return left / right
            raise ValueError("Operador nao permitido.")
        if isinstance(node, ast.UnaryOp):
            value = _eval(node.operand)
            if isinstance(node.op, ast.UAdd):
                return value
            if isinstance(node.op, ast.USub):
                return -value
            raise ValueError("Operador unario nao permitido.")
        if isinstance(node, ast.Name):
            if node.id not in variables:
                raise ValueError(f"Variavel invalida: {node.id}.")
            return Decimal(variables[node.id])
        if isinstance(node, ast.Constant):
            if isinstance(node.value, (int, float)):
                return Decimal(str(node.value))
            raise ValueError("Constante nao permitida.")
        raise ValueError("Expressao invalida.")

    return _eval(parsed)


def _normalize_rule_formula(target_column, raw_formula):
    formula = (raw_formula or "").strip()
    if not formula:
        return None, "Informe a formula."
    compact = re.sub(r"\s+", "", formula).upper()
    if "=" in compact:
        left, right = compact.split("=", 1)
        if left != target_column:
            return None, (
                f"O lado esquerdo deve ser {target_column}. "
                f"Exemplo: {target_column}=A/2*1.6"
            )
        compact = right
    if not compact:
        return None, "Formula invalida."
    return compact, None


def _resolve_bank_hours_formulas_for_month(competence_month):
    formulas = {"B": None, "C": None, "F": None}
    rules = (
        BankHoursRule.objects.filter(
            start_month__lte=competence_month,
            end_month__gte=competence_month,
        )
        .order_by("target_column", "-start_month", "-id")
    )
    latest_by_target = {}
    for rule in rules:
        latest_by_target.setdefault(rule.target_column, rule)
    for target_column, rule in latest_by_target.items():
        formulas[target_column] = rule.formula
    return formulas


def _calculate_bank_hours_columns(overtime_minutes, absence_minutes, manual_minutes, formulas):
    values = {
        "A": Decimal(overtime_minutes),
        "B": None,
        "C": None,
        "D": Decimal(absence_minutes),
        "E": Decimal(manual_minutes),
        "F": None,
    }
    missing_targets = []
    evaluation_errors = {}
    for target in ("B", "C", "F"):
        formula = formulas.get(target)
        if not formula:
            missing_targets.append(target)
            continue
        try:
            raw_result = _safe_decimal_formula_eval(formula, values)
            values[target] = raw_result.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        except Exception as error:
            missing_targets.append(target)
            evaluation_errors[target] = str(error)

    b_value = values["B"] if values["B"] is not None else Decimal(0)
    c_value = values["C"] if values["C"] is not None else Decimal(0)
    f_value = values["F"] if values["F"] is not None else Decimal(0)
    return {
        "overtime_minutes": int(values["A"]),
        "hour_bank_base_minutes": int(b_value),
        "hour_bank_bonus_minutes": int(c_value),
        "absence_minutes": int(values["D"]),
        "manual_minutes": int(values["E"]),
        "month_delta_minutes": int(f_value),
        "credit_minutes": int(b_value + c_value),
        "is_b_defined": values["B"] is not None,
        "is_c_defined": values["C"] is not None,
        "is_f_defined": values["F"] is not None,
        "missing_targets": missing_targets,
        "evaluation_errors": evaluation_errors,
        "formulas": formulas,
    }


def _count_weekday_occurrences_in_month(month_start):
    weekday_occurrences = [0] * 7
    _, days_in_month = monthrange(month_start.year, month_start.month)

    for day in range(1, days_in_month + 1):
        day_weekday = date(month_start.year, month_start.month, day).weekday()
        weekday_occurrences[day_weekday] += 1

    return weekday_occurrences


def _iter_clipped_dates_in_month(range_start, range_end, month_start, month_end):
    clipped_start = max(range_start, month_start)
    clipped_end = min(range_end, month_end)
    if clipped_end < clipped_start:
        return

    current_date = clipped_start
    while current_date <= clipped_end:
        yield current_date
        current_date += timedelta(days=1)


def _build_calendar_exception_dates_by_calendar(calendar_ids, month_start, month_end):
    if not calendar_ids:
        return {}

    periods = WorkCalendarPeriod.objects.filter(
        calendar_id__in=calendar_ids,
        start_date__lte=month_end,
        end_date__gte=month_start,
    )
    exception_dates_by_calendar = {}
    for period in periods:
        calendar_dates = exception_dates_by_calendar.setdefault(period.calendar_id, set())
        for period_day in _iter_clipped_dates_in_month(
            period.start_date,
            period.end_date,
            month_start,
            month_end,
        ):
            calendar_dates.add(period_day)

    return exception_dates_by_calendar


def _build_employee_exception_dates_by_employee(employee_ids, month_start, month_end):
    if not employee_ids:
        return {}

    exception_types = (
        EmployeeEvent.EVENT_TYPE_DAY_OFF,
        EmployeeEvent.EVENT_TYPE_VACATION,
        EmployeeEvent.EVENT_TYPE_ABSENCE,
        EmployeeEvent.EVENT_TYPE_MEDICAL_CERTIFICATE,
    )
    events = EmployeeEvent.objects.filter(
        employee_id__in=employee_ids,
        event_type__in=exception_types,
    ).filter(
        Q(end_date__isnull=True, effective_date__gte=month_start, effective_date__lte=month_end)
        | Q(end_date__isnull=False, effective_date__lte=month_end, end_date__gte=month_start)
    )

    exception_dates_by_employee = {}
    for event in events:
        event_end_date = event.end_date or event.effective_date
        employee_dates = exception_dates_by_employee.setdefault(event.employee_id, set())
        for event_day in _iter_clipped_dates_in_month(
            event.effective_date,
            event_end_date,
            month_start,
            month_end,
        ):
            employee_dates.add(event_day)

    return exception_dates_by_employee


def _build_employee_lifecycle_dates(employee_ids):
    lifecycle_dates = {
        employee_id: {"hiring_date": None, "termination_date": None}
        for employee_id in employee_ids
    }
    if not employee_ids:
        return lifecycle_dates

    events = (
        EmployeeEvent.objects.filter(
            employee_id__in=employee_ids,
            event_type__in=(
                EmployeeEvent.EVENT_TYPE_HIRING,
                EmployeeEvent.EVENT_TYPE_TERMINATION,
            ),
        )
        .only("employee_id", "event_type", "effective_date")
        .order_by("employee_id", "effective_date", "id")
    )
    for event in events:
        employee_lifecycle = lifecycle_dates.setdefault(
            event.employee_id,
            {"hiring_date": None, "termination_date": None},
        )
        if event.event_type == EmployeeEvent.EVENT_TYPE_HIRING:
            if (
                employee_lifecycle["hiring_date"] is None
                or event.effective_date < employee_lifecycle["hiring_date"]
            ):
                employee_lifecycle["hiring_date"] = event.effective_date
            continue

        if (
            employee_lifecycle["termination_date"] is None
            or event.effective_date < employee_lifecycle["termination_date"]
        ):
            employee_lifecycle["termination_date"] = event.effective_date

    return lifecycle_dates


def _get_employee_active_date_range(employee_id, range_start, range_end, lifecycle_dates):
    lifecycle = lifecycle_dates.get(employee_id, {})
    hiring_date = lifecycle.get("hiring_date")
    termination_date = lifecycle.get("termination_date")

    active_start = max(range_start, hiring_date) if hiring_date else range_start
    active_end = min(range_end, termination_date) if termination_date else range_end
    if active_end < active_start:
        return None, None

    return active_start, active_end


def _filter_employees_active_in_period(employees, range_start, range_end, lifecycle_dates):
    return [
        employee
        for employee in employees
        if _get_employee_active_date_range(
            employee.id,
            range_start,
            range_end,
            lifecycle_dates,
        )
        != (None, None)
    ]


def _build_employee_schedule_events_by_employee(employee_ids):
    if not employee_ids:
        return {}

    events = (
        EmployeeEvent.objects.select_related(
            "previous_work_schedule__calendar",
            "new_work_schedule__calendar",
        )
        .filter(
            employee_id__in=employee_ids,
            event_type=EmployeeEvent.EVENT_TYPE_ALLOCATION_CHANGE,
        )
        .order_by("employee_id", "effective_date", "id")
    )
    schedule_events_by_employee = defaultdict(list)
    for event in events:
        if not event.previous_work_schedule_id and not event.new_work_schedule_id:
            continue
        schedule_events_by_employee[event.employee_id].append(event)

    return schedule_events_by_employee


def _resolve_work_schedule_for_employee_date(
    employee,
    target_date,
    schedule_events_by_employee,
):
    schedule_events = schedule_events_by_employee.get(employee.id, ())
    if not schedule_events:
        return employee.work_schedule

    latest_past_event = None
    earliest_future_event = None
    for event in schedule_events:
        if event.effective_date <= target_date:
            latest_past_event = event
            continue
        earliest_future_event = event
        break

    if latest_past_event:
        return latest_past_event.new_work_schedule
    if earliest_future_event:
        return earliest_future_event.previous_work_schedule

    return employee.work_schedule


def _collect_calendar_ids_for_timesheet(employees, schedule_events_by_employee):
    calendar_ids = {
        employee.work_schedule.calendar_id
        for employee in employees
        if employee.work_schedule and employee.work_schedule.calendar_id
    }
    for schedule_events in schedule_events_by_employee.values():
        for event in schedule_events:
            for schedule in (event.previous_work_schedule, event.new_work_schedule):
                if schedule and schedule.calendar_id:
                    calendar_ids.add(schedule.calendar_id)

    return calendar_ids


def _collect_work_schedules_used_for_period(
    employees,
    range_start,
    range_end,
    lifecycle_dates,
    schedule_events_by_employee,
):
    schedules_by_id = {}
    for employee in employees:
        active_start, active_end = _get_employee_active_date_range(
            employee.id,
            range_start,
            range_end,
            lifecycle_dates,
        )
        if active_start is None:
            continue

        current_date = active_start
        while current_date <= active_end:
            schedule = _resolve_work_schedule_for_employee_date(
                employee,
                current_date,
                schedule_events_by_employee,
            )
            if schedule:
                schedules_by_id[schedule.id] = schedule
            current_date += timedelta(days=1)

    return schedules_by_id


def _get_timesheet_base_employees():
    return list(
        Employee.objects.select_related(
            "sector",
            "sector__department",
            "sector__department__area",
            "work_schedule",
            "work_schedule__calendar",
        )
        .filter(deactivated_at__isnull=True)
        .order_by("nome_completo")
    )


def _calculate_expected_minutes_for_employee(
    employee,
    month_start,
    month_end,
    calendar_exception_dates_by_calendar,
    employee_exception_dates_by_employee,
    schedule_events_by_employee=None,
    lifecycle_dates=None,
):
    schedule_events_by_employee = schedule_events_by_employee or {}
    lifecycle_dates = lifecycle_dates or {}
    active_start, active_end = _get_employee_active_date_range(
        employee.id,
        month_start,
        month_end,
        lifecycle_dates,
    )
    if active_start is None:
        return 0

    employee_exception_dates = employee_exception_dates_by_employee.get(employee.id, set())

    total_hours = Decimal("0")
    current_date = active_start
    while current_date <= active_end:
        work_schedule = _resolve_work_schedule_for_employee_date(
            employee,
            current_date,
            schedule_events_by_employee,
        )
        calendar_exception_dates = (
            calendar_exception_dates_by_calendar.get(work_schedule.calendar_id, set())
            if work_schedule and work_schedule.calendar_id
            else set()
        )
        if (
            work_schedule
            and current_date not in calendar_exception_dates
            and current_date not in employee_exception_dates
        ):
            field_name = SCHEDULE_FIELDS_IN_WEEKDAY_ORDER[current_date.weekday()]
            day_hours = getattr(work_schedule, field_name, Decimal("0")) or Decimal("0")
            total_hours += day_hours
        current_date += timedelta(days=1)

    return int(
        (total_hours * Decimal("60")).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


def _get_work_schedule_or_error(work_schedule_id):
    if not work_schedule_id:
        return None, None

    work_schedule = WorkSchedule.objects.filter(id=work_schedule_id).first()
    if not work_schedule:
        return None, "Selecione uma escala valida."

    return work_schedule, None


def _get_work_calendar_or_error(calendar_id):
    if not calendar_id:
        return None, "Selecione um calendario valido."

    calendar = WorkCalendar.objects.filter(id=calendar_id).first()
    if not calendar:
        return None, "Selecione um calendario valido."

    return calendar, None


def _parse_employee_characteristics(request_data):
    tipo = (request_data.get("tipo") or "").strip()
    regime_compensacao_jornada = (
        request_data.get("regime_compensacao_jornada") or ""
    ).strip()
    valid_types = {choice[0] for choice in Employee.TYPE_CHOICES}
    valid_compensation_regimes = {
        choice[0] for choice in Employee.REGIME_COMPENSACAO_JORNADA_CHOICES
    }

    if tipo not in valid_types:
        return None, None, "Selecione um tipo de empregado valido."

    if regime_compensacao_jornada not in valid_compensation_regimes:
        return (
            None,
            None,
            "Selecione um regime de compensacao de jornada valido.",
        )

    return tipo, regime_compensacao_jornada, None


def _normalize_csv_header(value):
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return normalized.strip().lower()


def _normalize_employee_type(value):
    normalized = _normalize_csv_header(value)
    if normalized in {"direto", "direct"}:
        return Employee.TYPE_DIRETO
    if normalized in {"indireto", "indirect"}:
        return Employee.TYPE_INDIRETO
    return None


def _normalize_name_key(value):
    return " ".join((value or "").strip().lower().split())


def _get_or_create_default_area_department():
    area_zero = Area.objects.filter(nome="0", deactivated_at__isnull=True).first()
    if not area_zero:
        area_zero = Area.objects.create(nome="0")

    department_zero = Department.objects.filter(
        area=area_zero,
        nome="0",
        deactivated_at__isnull=True,
    ).first()
    if not department_zero:
        department_zero = Department.objects.create(area=area_zero, nome="0")

    return area_zero, department_zero


def _parse_br_date(raw_value):
    normalized = (raw_value or "").strip()
    try:
        day, month, year = normalized.split("/")
        return date(int(year), int(month), int(day))
    except Exception as error:
        raise ValueError(str(error))


def _import_employees_from_csv(uploaded_file):
    if not uploaded_file:
        return 0, ["Selecione um arquivo CSV para importar."]

    try:
        raw_content = uploaded_file.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        return 0, ["Nao foi possivel ler o CSV em UTF-8."]

    sample = raw_content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(io.StringIO(raw_content), dialect=dialect)
    if not reader.fieldnames:
        return 0, ["Arquivo CSV sem cabecalho."]

    field_map = {_normalize_csv_header(name): name for name in reader.fieldnames}
    required_headers = ("matricula", "admissao", "nome", "setor", "cargo", "tipo")
    missing_headers = [header for header in required_headers if header not in field_map]
    if missing_headers:
        return 0, [
            "CSV invalido: faltam colunas obrigatorias "
            f"({', '.join(missing_headers)})."
        ]

    imported = 0
    errors = []
    seen_csv_ids = {}
    seen_csv_name_date = {}
    existing_ids = set(
        Employee.objects.exclude(matricula__isnull=True)
        .exclude(matricula="")
        .values_list("matricula", flat=True)
    )
    existing_name_date = {
        (_normalize_name_key(name), effective_date)
        for name, effective_date in EmployeeEvent.objects.filter(
            event_type=EmployeeEvent.EVENT_TYPE_HIRING
        ).values_list("employee__nome_completo", "effective_date")
    }

    for line_index, row in enumerate(reader, start=2):
        matricula = (row.get(field_map["matricula"]) or "").strip()
        admission_raw = (row.get(field_map["admissao"]) or "").strip()
        nome = (row.get(field_map["nome"]) or "").strip()
        sector_name = (row.get(field_map["setor"]) or "").strip()
        cargo_name = (row.get(field_map["cargo"]) or "").strip()
        tipo_raw = (row.get(field_map["tipo"]) or "").strip()
        tipo = _normalize_employee_type(tipo_raw)

        if not nome:
            errors.append(f"Linha {line_index}: nome obrigatorio.")
            continue
        if tipo == Employee.TYPE_DIRETO and not matricula:
            errors.append(f"Linha {line_index}: ID/matricula obrigatorio para empregado direto.")
            continue
        if not sector_name:
            errors.append(f"Linha {line_index}: setor obrigatorio.")
            continue
        if not cargo_name:
            errors.append(f"Linha {line_index}: cargo obrigatorio.")
            continue
        if tipo is None:
            errors.append(
                f"Linha {line_index}: tipo invalido '{tipo_raw}'. Use direto ou indireto."
            )
            continue
        try:
            admission_date = _parse_br_date(admission_raw)
        except ValueError:
            errors.append(
                f"Linha {line_index}: data de admissao invalida '{admission_raw}'. Use DD/MM/AAAA."
            )
            continue

        if tipo == Employee.TYPE_DIRETO:
            first_line_for_id = seen_csv_ids.get(matricula)
            if first_line_for_id:
                errors.append(
                    f"Linha {line_index}: ID/matricula '{matricula}' repetido no CSV (primeira ocorrencia na linha {first_line_for_id})."
                )
                continue
            if matricula in existing_ids:
                errors.append(
                    f"Linha {line_index}: ID/matricula '{matricula}' ja existe no banco."
                )
                continue

        name_date_key = (_normalize_name_key(nome), admission_date)
        first_line_for_name_date = seen_csv_name_date.get(name_date_key)
        if first_line_for_name_date:
            errors.append(
                f"Linha {line_index}: nome '{nome}' com mesma data de admissao ({admission_date.isoformat()}) repetido no CSV."
            )
            continue
        if name_date_key in existing_name_date:
            errors.append(
                f"Linha {line_index}: ja existe empregado com nome '{nome}' e data de admissao {admission_date.isoformat()}."
            )
            continue

        cargo = Cargo.objects.filter(
            nome__iexact=cargo_name,
            deactivated_at__isnull=True,
        ).first()
        if not cargo:
            cargo = Cargo.objects.create(nome=cargo_name)

        sector = Sector.objects.filter(
            nome__iexact=sector_name,
            deactivated_at__isnull=True,
        ).order_by("id").first()
        if not sector:
            _, department_zero = _get_or_create_default_area_department()
            sector = Sector.objects.create(
                nome=sector_name,
                department=department_zero,
            )

        with transaction.atomic():
            employee = Employee.objects.create(
                matricula=matricula if tipo == Employee.TYPE_DIRETO else None,
                nome_completo=nome,
                tipo=tipo,
                regime_compensacao_jornada=Employee.REGIME_COMPENSACAO_NAO_PARTICIPANTE,
                cargo=cargo,
                sector=sector,
            )
            if tipo == Employee.TYPE_INDIRETO:
                employee.matricula = str(employee.id + 100000)
                employee.save(update_fields=["matricula"])

            EmployeeEvent.objects.create(
                employee=employee,
                event_type=EmployeeEvent.EVENT_TYPE_HIRING,
                effective_date=admission_date,
                notes="Evento de admissao criado automaticamente por importacao CSV.",
            )
        if tipo == Employee.TYPE_DIRETO:
            seen_csv_ids[matricula] = line_index
        seen_csv_name_date[name_date_key] = line_index
        if tipo == Employee.TYPE_DIRETO:
            existing_ids.add(matricula)
        existing_name_date.add(name_date_key)
        imported += 1

    return imported, errors


def _active_taxonomy():
    areas = Area.objects.filter(deactivated_at__isnull=True).order_by("nome")
    departments = Department.objects.filter(
        deactivated_at__isnull=True,
        area__deactivated_at__isnull=True,
    ).select_related("area").order_by("area__nome", "nome")
    return areas, departments


def _parse_sector_taxonomy_or_error(request_data):
    area_id = (request_data.get("area_id") or "").strip()
    department_id = (request_data.get("department_id") or "").strip()
    sector_name = (request_data.get("nome") or "").strip()

    if not area_id:
        return None, None, None, "Selecione uma area."
    if not department_id:
        return None, None, None, "Selecione um departamento."
    if not sector_name:
        return None, None, None, "Nome do setor e obrigatorio."

    area = Area.objects.filter(id=area_id, deactivated_at__isnull=True).first()
    if not area:
        return None, None, None, "Selecione uma area valida."

    department = Department.objects.filter(
        id=department_id,
        area=area,
        deactivated_at__isnull=True,
    ).first()
    if not department:
        return None, None, None, "Selecione um departamento valido para a area."

    return area, department, sector_name, None


def employees_page(request):
    try:
        active_sectors = (
            Sector.objects.filter(
                deactivated_at__isnull=True,
            )
            .select_related("department__area")
            .order_by("department__area__nome", "department__nome", "nome")
        )
        active_cargos = Cargo.objects.filter(deactivated_at__isnull=True).order_by("nome")
        work_schedules = WorkSchedule.objects.all().order_by("nome")
        all_sectors = (
            Sector.objects.all()
            .select_related("department__area")
            .order_by("deactivated_at", "department__area__nome", "department__nome", "nome")
        )
        active_areas, active_departments = _active_taxonomy()
        import_errors = request.session.pop("employee_import_errors", None)
        import_result_count = request.session.pop("employee_import_result_count", None)

        if request.method == "POST":
            if (request.POST.get("form_type") or "").strip() == "employee_import_csv":
                imported_count, errors = _import_employees_from_csv(request.FILES.get("csv_file"))
                if errors:
                    request.session["employee_import_errors"] = errors
                    request.session["employee_import_result_count"] = imported_count
                elif imported_count:
                    messages.success(
                        request,
                        f"Importacao concluida: {imported_count} empregado(s) importado(s).",
                    )
                return redirect("employees_page")

            matricula = (request.POST.get("matricula") or "").strip()
            nome_completo = (request.POST.get("nome_completo") or "").strip()
            tipo, regime_compensacao_jornada, characteristics_error = (
                _parse_employee_characteristics(request.POST)
            )
            cargo_id = (request.POST.get("cargo_id") or "").strip()
            sector_id = (request.POST.get("sector_id") or "").strip()
            work_schedule_id = (request.POST.get("work_schedule_id") or "").strip()

            if not nome_completo or not cargo_id or not sector_id:
                messages.error(
                    request,
                    "Nome completo, cargo e setor sao obrigatorios.",
                )
                return redirect("employees_page")

            if characteristics_error:
                messages.error(request, characteristics_error)
                return redirect("employees_page")

            if tipo == Employee.TYPE_DIRETO and not matricula:
                messages.error(request, "Matricula e obrigatoria para empregado direto.")
                return redirect("employees_page")

            if matricula and Employee.objects.filter(matricula=matricula).exists():
                messages.error(request, "Ja existe empregado com essa matricula.")
                return redirect("employees_page")

            cargo = active_cargos.filter(id=cargo_id).first()
            if not cargo:
                messages.error(request, "Selecione um cargo ativo valido.")
                return redirect("employees_page")

            sector = active_sectors.filter(id=sector_id).first()
            if not sector:
                messages.error(request, "Selecione um setor ativo valido.")
                return redirect("employees_page")

            work_schedule, schedule_error = _get_work_schedule_or_error(work_schedule_id)
            if schedule_error:
                messages.error(request, schedule_error)
                return redirect("employees_page")

            Employee.objects.create(
                matricula=matricula or None,
                nome_completo=nome_completo,
                tipo=tipo,
                regime_compensacao_jornada=regime_compensacao_jornada,
                cargo=cargo,
                sector=sector,
                work_schedule=work_schedule,
            )

            messages.success(request, "Empregado cadastrado com sucesso.")
            return redirect("employees_page")

        employees = Employee.objects.select_related(
            "cargo",
            "sector__department__area",
            "work_schedule",
        ).order_by(
            "deactivated_at",
            "nome_completo",
        )
        active_count = employees.filter(deactivated_at__isnull=True).count()
        inactive_count = employees.filter(deactivated_at__isnull=False).count()
        participants_count = employees.filter(
            regime_compensacao_jornada=Employee.REGIME_COMPENSACAO_PARTICIPANTE
        ).count()
        unassigned_count = employees.filter(sector__isnull=True).count()
        selected_sector_id = (request.GET.get("selected_sector") or "").strip()

        return render(
            request,
            "employees.html",
            {
                "employees": employees,
                "employees_count": employees.count(),
                "active_count": active_count,
                "inactive_count": inactive_count,
                "participants_count": participants_count,
                "active_sectors": active_sectors,
                "active_cargos": active_cargos,
                "work_schedules": work_schedules,
                "all_sectors": all_sectors,
                "unassigned_count": unassigned_count,
                "employee_type_choices": Employee.TYPE_CHOICES,
                "compensation_regime_choices": Employee.REGIME_COMPENSACAO_JORNADA_CHOICES,
                "open_employee_modal": request.GET.get("open_employee_modal") == "1",
                "open_sector_modal": request.GET.get("open_sector_modal") == "1",
                "selected_sector_id": selected_sector_id,
                "active_areas": active_areas,
                "active_departments": active_departments,
                "employee_import_errors": import_errors or [],
                "employee_import_result_count": import_result_count,
            },
        )
    except (OperationalError, ProgrammingError):
        messages.error(
            request,
            (
                "Banco de dados indisponivel ou sem estrutura inicial. "
                "Vá em Configuracoes de banco para criar um novo .db do zero "
                "ou informar o caminho de um banco existente."
            ),
        )
        return redirect("database_settings_page")


def _iterate_competence_months(start_date, end_date):
    current_month = date(start_date.year, start_date.month, 1)
    end_month = date(end_date.year, end_date.month, 1)

    while current_month <= end_month:
        yield current_month
        if current_month.month == 12:
            current_month = date(current_month.year + 1, 1, 1)
        else:
            current_month = date(current_month.year, current_month.month + 1, 1)


def _find_first_closed_competence_month(start_date, end_date=None):
    range_end = end_date or start_date
    competence_months = list(_iterate_competence_months(start_date, range_end))
    return (
        TimesheetMonthClosure.objects.filter(competence_month__in=competence_months)
        .order_by("competence_month")
        .values_list("competence_month", flat=True)
        .first()
    )


def events_page(request):
    employees = Employee.objects.select_related("sector").order_by(
        "deactivated_at",
        "nome_completo",
    )
    event_type_choices = EmployeeEvent.EVENT_TYPE_CHOICES
    valid_event_types = {choice[0] for choice in event_type_choices}

    if request.method == "POST":
        employee_id = (request.POST.get("employee_id") or "").strip()
        event_type = (request.POST.get("event_type") or "").strip()
        raw_effective_date = (request.POST.get("effective_date") or "").strip()
        raw_end_date = (request.POST.get("end_date") or "").strip()
        raw_bank_hours_amount = (request.POST.get("bank_hours_amount") or "").strip()
        notes = (request.POST.get("notes") or "").strip()

        employee = Employee.objects.filter(id=employee_id).first()
        if not employee:
            messages.error(request, "Selecione um empregado valido.")
            return redirect(_events_redirect_with_filters())

        if event_type not in valid_event_types:
            messages.error(request, "Selecione um tipo de evento valido.")
            return redirect(_events_redirect_with_filters(employee_id=employee.id))

        if not raw_effective_date:
            messages.error(request, "Informe a data de inicio do evento.")
            return redirect(
                _events_redirect_with_filters(
                    employee_id=employee.id,
                    event_type=event_type,
                )
            )

        try:
            effective_date = date.fromisoformat(raw_effective_date)
        except ValueError:
            messages.error(request, "Informe uma data de inicio valida.")
            return redirect(
                _events_redirect_with_filters(
                    employee_id=employee.id,
                    event_type=event_type,
                )
            )

        end_date = None
        if raw_end_date:
            try:
                end_date = date.fromisoformat(raw_end_date)
            except ValueError:
                messages.error(request, "Informe uma data final valida.")
                return redirect(
                    _events_redirect_with_filters(
                        employee_id=employee.id,
                        event_type=event_type,
                    )
                )

            if end_date < effective_date:
                messages.error(
                    request,
                    "A data final nao pode ser menor que a data de inicio.",
                )
                return redirect(
                    _events_redirect_with_filters(
                        employee_id=employee.id,
                        event_type=event_type,
                    )
                )

        first_closed_month = _find_first_closed_competence_month(
            effective_date,
            end_date or effective_date,
        )
        if first_closed_month:
            messages.error(
                request,
                "Nao e permitido registrar evento/ajuste com data em competencia encerrada "
                f"({_format_competence_month_label(first_closed_month)}).",
            )
            return redirect(
                _events_redirect_with_filters(
                    employee_id=employee.id,
                    event_type=event_type,
                )
            )
        bank_hours_amount = None
        if raw_bank_hours_amount:
            try:
                bank_hours_amount = Decimal(raw_bank_hours_amount.replace(",", "."))
            except InvalidOperation:
                messages.error(request, "Informe um valor valido para banco de horas.")
                return redirect(
                    _events_redirect_with_filters(
                        employee_id=employee.id,
                        event_type=event_type,
                    )
                )

        if event_type == EmployeeEvent.EVENT_TYPE_BANK_HOURS_MOVEMENT:
            if bank_hours_amount is None or bank_hours_amount == 0:
                messages.error(
                    request,
                    "Informe o valor em horas para movimento de banco de horas.",
                )
                return redirect(
                    _events_redirect_with_filters(
                        employee_id=employee.id,
                        event_type=event_type,
                    )
                )
        else:
            bank_hours_amount = None

        EmployeeEvent.objects.create(
            employee=employee,
            event_type=event_type,
            effective_date=effective_date,
            end_date=end_date,
            bank_hours_amount=bank_hours_amount,
            notes=notes,
        )

        if event_type == EmployeeEvent.EVENT_TYPE_BANK_HOURS_ADOPTION:
            if (
                employee.regime_compensacao_jornada
                != Employee.REGIME_COMPENSACAO_PARTICIPANTE
            ):
                employee.regime_compensacao_jornada = (
                    Employee.REGIME_COMPENSACAO_PARTICIPANTE
                )
                employee.save(update_fields=["regime_compensacao_jornada"])
        elif event_type == EmployeeEvent.EVENT_TYPE_BANK_HOURS_WITHDRAWAL:
            if (
                employee.regime_compensacao_jornada
                != Employee.REGIME_COMPENSACAO_NAO_PARTICIPANTE
            ):
                employee.regime_compensacao_jornada = (
                    Employee.REGIME_COMPENSACAO_NAO_PARTICIPANTE
                )
                employee.save(update_fields=["regime_compensacao_jornada"])

        messages.success(request, "Evento registrado com sucesso.")
        return redirect(
            _events_redirect_with_filters(
                employee_id=employee.id,
                event_type=event_type,
            )
        )

    selected_employee = None
    employee_id = (request.GET.get("employee_id") or "").strip()
    selected_event_type = (request.GET.get("event_type") or "").strip()

    if employee_id.isdigit():
        selected_employee = employees.filter(id=employee_id).first()

    if selected_event_type not in valid_event_types:
        selected_event_type = ""

    events = EmployeeEvent.objects.select_related(
        "employee",
        "previous_sector",
        "new_sector",
        "previous_work_schedule",
        "new_work_schedule",
    )
    if selected_employee:
        events = events.filter(employee=selected_employee)
    if selected_event_type:
        events = events.filter(event_type=selected_event_type)

    return render(
        request,
        "events.html",
        {
            "employees": employees,
            "selected_employee": selected_employee,
            "selected_event_type": selected_event_type,
            "event_type_choices": event_type_choices,
            "events": events,
            "events_count": events.count(),
        },
    )


def _format_competence_month_label(competence_month):
    return competence_month.strftime("%m/%Y")


def _format_minutes_as_hour_label(total_minutes):
    hours = total_minutes // 60
    minutes = total_minutes % 60
    return f"{hours:02d}:{minutes:02d}"


def _format_signed_minutes_as_hour_label(total_minutes):
    sign = "-" if total_minutes < 0 else ""
    absolute_minutes = abs(total_minutes)
    hours = absolute_minutes // 60
    minutes = absolute_minutes % 60
    return f"{sign}{hours:02d}:{minutes:02d}"


def _get_month_date_range(month_start):
    _, days_in_month = monthrange(month_start.year, month_start.month)
    month_end = date(month_start.year, month_start.month, days_in_month)
    return month_start, month_end


def _month_start_for_date(value):
    return date(value.year, value.month, 1)


def _minutes_from_decimal_hours(raw_hours):
    decimal_hours = Decimal(raw_hours)
    return int((decimal_hours * Decimal("60")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _decimal_hours_from_minutes(total_minutes):
    return (
        Decimal(total_minutes) / Decimal("60")
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _resolve_semester_window(year, semester):
    semester_months = (1, 2, 3, 4, 5, 6) if semester == 1 else (7, 8, 9, 10, 11, 12)
    competence_months = [date(year, month, 1) for month in semester_months]
    start_month = competence_months[0]
    end_month = competence_months[-1]
    adjustment_date = date(year, semester_months[-1], monthrange(year, semester_months[-1])[1])
    return semester_months, competence_months, start_month, end_month, adjustment_date


def _is_month_locked_by_bank_hours_semester_closure(competence_month):
    semester = 1 if competence_month.month <= 6 else 2
    return BankHoursSemesterClosure.objects.filter(
        year=competence_month.year,
        semester=semester,
        reversed_at__isnull=True,
    ).exists()


def _calc_bank_credit_minutes(overtime_minutes):
    return int(
        (
            Decimal(overtime_minutes)
            * Decimal("0.5")
            * Decimal("1.6")
        ).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    )


def _create_month_closure_snapshots(closure):
    month_start, month_end = _get_month_date_range(closure.competence_month)

    base_employees = _get_timesheet_base_employees()
    base_employee_ids = [employee.id for employee in base_employees]
    lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
    active_employees = _filter_employees_active_in_period(
        base_employees,
        month_start,
        month_end,
        lifecycle_dates,
    )
    employee_ids = [employee.id for employee in active_employees]
    schedule_events_by_employee = _build_employee_schedule_events_by_employee(employee_ids)

    schedules_by_id = _collect_work_schedules_used_for_period(
        active_employees,
        month_start,
        month_end,
        lifecycle_dates,
        schedule_events_by_employee,
    )
    calendar_names_by_id = {}
    for schedule in schedules_by_id.values():
        if schedule.calendar_id:
            calendar_names_by_id[schedule.calendar_id] = schedule.calendar.nome

    schedule_snapshots = []
    for schedule in schedules_by_id.values():
        calendar = schedule.calendar
        schedule_snapshots.append(
            TimesheetMonthClosureWorkScheduleSnapshot(
                closure=closure,
                source_work_schedule_id=schedule.id,
                work_schedule_name=schedule.nome,
                source_calendar_id=calendar.id if calendar else None,
                calendar_name=calendar.nome if calendar else "",
                horas_segunda=schedule.horas_segunda,
                horas_terca=schedule.horas_terca,
                horas_quarta=schedule.horas_quarta,
                horas_quinta=schedule.horas_quinta,
                horas_sexta=schedule.horas_sexta,
                horas_sabado=schedule.horas_sabado,
                horas_domingo=schedule.horas_domingo,
            )
        )

    if schedule_snapshots:
        TimesheetMonthClosureWorkScheduleSnapshot.objects.bulk_create(schedule_snapshots)

    if not calendar_names_by_id:
        return

    period_snapshots = []
    periods = WorkCalendarPeriod.objects.filter(
        calendar_id__in=calendar_names_by_id.keys(),
        start_date__lte=month_end,
        end_date__gte=month_start,
    ).order_by("calendar_id", "start_date", "id")

    for period in periods:
        period_snapshots.append(
            TimesheetMonthClosureCalendarPeriodSnapshot(
                closure=closure,
                source_calendar_id=period.calendar_id,
                calendar_name=calendar_names_by_id.get(period.calendar_id, ""),
                period_type=period.period_type,
                start_date=period.start_date,
                end_date=period.end_date,
                description=period.description,
            )
        )

    if period_snapshots:
        TimesheetMonthClosureCalendarPeriodSnapshot.objects.bulk_create(period_snapshots)

    individual_exception_types = (
        EmployeeEvent.EVENT_TYPE_VACATION,
        EmployeeEvent.EVENT_TYPE_DAY_OFF,
        EmployeeEvent.EVENT_TYPE_ABSENCE,
        EmployeeEvent.EVENT_TYPE_MEDICAL_CERTIFICATE,
    )
    event_snapshots = []
    exception_events = (
        EmployeeEvent.objects.select_related("employee")
        .filter(
            event_type__in=individual_exception_types,
            effective_date__lte=month_end,
        )
        .filter(
            Q(end_date__gte=month_start)
            | Q(end_date__isnull=True, effective_date__gte=month_start)
        )
        .order_by("employee__nome_completo", "effective_date", "id")
    )
    for event in exception_events:
        event_snapshots.append(
            TimesheetMonthClosureEmployeeEventSnapshot(
                closure=closure,
                source_employee_event_id=event.id,
                source_employee_id=event.employee_id,
                employee_registration=event.employee.matricula,
                employee_name=event.employee.nome_completo,
                event_type=event.event_type,
                effective_date=event.effective_date,
                end_date=event.end_date,
                notes=event.notes,
            )
        )

    if event_snapshots:
        TimesheetMonthClosureEmployeeEventSnapshot.objects.bulk_create(event_snapshots)


def _delete_month_closures_with_snapshots(closure_qs):
    closure_ids = list(closure_qs.values_list("id", flat=True))
    if not closure_ids:
        return 0

    TimesheetMonthClosureWorkScheduleSnapshot.objects.filter(
        closure_id__in=closure_ids
    )._raw_delete(closure_qs.db)
    TimesheetMonthClosureCalendarPeriodSnapshot.objects.filter(
        closure_id__in=closure_ids
    )._raw_delete(closure_qs.db)
    TimesheetMonthClosureEmployeeEventSnapshot.objects.filter(
        closure_id__in=closure_ids
    )._raw_delete(closure_qs.db)
    return TimesheetMonthClosure.objects.filter(id__in=closure_ids)._raw_delete(
        closure_qs.db
    )


def _get_latest_closure_for_schedule(schedule_id):
    schedule_snapshot = (
        TimesheetMonthClosureWorkScheduleSnapshot.objects.select_related("closure")
        .filter(source_work_schedule_id=schedule_id)
        .order_by("-closure__competence_month")
        .first()
    )
    return schedule_snapshot.closure if schedule_snapshot else None


def _get_latest_closure_for_calendar(calendar_id):
    schedule_hit = (
        TimesheetMonthClosureWorkScheduleSnapshot.objects.select_related("closure")
        .filter(source_calendar_id=calendar_id)
        .order_by("-closure__competence_month")
        .first()
    )
    period_hit = (
        TimesheetMonthClosureCalendarPeriodSnapshot.objects.select_related("closure")
        .filter(source_calendar_id=calendar_id)
        .order_by("-closure__competence_month")
        .first()
    )

    candidates = [hit.closure for hit in (schedule_hit, period_hit) if hit]
    if not candidates:
        return None

    return max(candidates, key=lambda closure: closure.competence_month)


def timesheet_page(request):
    if request.method == "POST":
        competence_month, competence_month_value, error = _resolve_competence_month(
            request.POST.get("competence_month")
        )
        if error:
            messages.error(request, error)
            return redirect(_timesheet_redirect_with_filters())

        action = (request.POST.get("action") or "save").strip()
        competence_month_label = _format_competence_month_label(competence_month)
        month_start, month_end = _get_month_date_range(competence_month)
        base_employees = _get_timesheet_base_employees()
        base_employee_ids = [employee.id for employee in base_employees]
        lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
        employees = _filter_employees_active_in_period(
            base_employees,
            month_start,
            month_end,
            lifecycle_dates,
        )

        if action == "close_month":
            closure, created = TimesheetMonthClosure.objects.get_or_create(
                competence_month=competence_month
            )
            if created:
                _create_month_closure_snapshots(closure)
                messages.success(
                    request,
                    f"Timesheet de {competence_month_label} encerrado. Alteracoes estao bloqueadas.",
                )
            else:
                messages.info(
                    request,
                    f"Timesheet de {competence_month_label} ja estava encerrado.",
                )
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        if action == "reopen_month":
            closure_qs = TimesheetMonthClosure.objects.filter(
                competence_month=competence_month
            )
            was_closed = closure_qs.exists()
            if was_closed:
                try:
                    _delete_month_closures_with_snapshots(closure_qs)
                except IntegrityError:
                    messages.error(
                        request,
                        (
                            f"Nao foi possivel reabrir o timesheet de {competence_month_label} "
                            "porque existem dependencias no banco. "
                            "Revise referencias de snapshot para esta competencia."
                        ),
                    )
                    return redirect(
                        _timesheet_redirect_with_filters(
                            competence_month=competence_month_value
                        )
                    )
                messages.success(
                    request,
                    f"Timesheet de {competence_month_label} reaberto. Alteracoes estao liberadas.",
                )
            else:
                messages.info(
                    request,
                    f"Timesheet de {competence_month_label} ja estava aberto.",
                )
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        if action == "preview_import_timesheet":
            preview, hard_errors, warnings = _build_timesheet_import_preview(
                request.FILES.get("timesheet_import_file"),
                competence_month,
            )
            report = None
            if preview:
                with transaction.atomic():
                    report = _create_timesheet_import_report(preview, competence_month)
            if hard_errors:
                messages.error(request, "Relatorio de importacao gerado com erro: " + " | ".join(hard_errors[:3]))
            elif warnings or (preview and (preview["skipped_rows"] or preview["missing_employee_rows"])):
                messages.warning(request, "Arquivo lido com alertas. Revise o relatorio antes de aprovar.")
            else:
                messages.success(request, "Arquivo validado. Revise e aprove o relatorio para preencher o timesheet.")
            if report:
                return redirect(
                    reverse(
                        "timesheet_import_report_page",
                        kwargs={"report_id": report.id},
                    )
                )
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        if action == "confirm_import_timesheet":
            messages.error(request, "A importacao deve ser aprovada pela subpagina do relatorio.")
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )
        if TimesheetMonthClosure.objects.filter(competence_month=competence_month).exists():
            messages.error(
                request,
                f"O timesheet de {competence_month_label} esta encerrado. Reabra o mes para salvar alteracoes.",
            )
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        entries_in_month = {
            entry.employee_id: entry
            for entry in EmployeeTimeEntry.objects.filter(
                competence_month=competence_month
            ).only("id", "employee_id")
        }

        updated_count = 0
        removed_count = 0
        is_bank_hours_column_locked = _is_month_locked_by_bank_hours_semester_closure(
            competence_month
        )
        for employee in employees:
            parsed_values = {}
            for field_name, field_label in TIME_ENTRY_MINUTE_FIELDS:
                value, value_error = _parse_non_negative_minutes(
                    request.POST.get(f"{field_name}_{employee.id}"),
                    f"{field_label} - {employee.nome_completo}",
                )
                if value_error:
                    messages.error(request, value_error)
                    return redirect(
                        _timesheet_redirect_with_filters(
                            competence_month=competence_month_value
                        )
                    )
                parsed_values[field_name] = value

            if (
                employee.regime_compensacao_jornada
                != Employee.REGIME_COMPENSACAO_PARTICIPANTE
            ):
                parsed_values["absence_bank_minutes"] = 0

            notes = (request.POST.get(f"notes_{employee.id}") or "").strip()
            has_content = any(value > 0 for value in parsed_values.values()) or notes
            existing_entry = entries_in_month.get(employee.id)

            if (
                is_bank_hours_column_locked
                and employee.regime_compensacao_jornada == Employee.REGIME_COMPENSACAO_PARTICIPANTE
            ):
                parsed_values["absence_bank_minutes"] = (
                    existing_entry.absence_bank_minutes if existing_entry else 0
                )

            if has_content:
                EmployeeTimeEntry.objects.update_or_create(
                    employee=employee,
                    competence_month=competence_month,
                    defaults={
                        **parsed_values,
                        "notes": notes,
                    },
                )
                updated_count += 1
                continue

            if existing_entry:
                existing_entry.delete()
                removed_count += 1

        messages.success(
            request,
            f"Timesheet mensal salvo para {updated_count} empregado(s).",
        )
        if removed_count:
            messages.info(
                request,
                f"{removed_count} empregado(s) sem valores foram removidos do mes.",
            )

        return redirect(
            _timesheet_redirect_with_filters(competence_month=competence_month_value)
        )

    competence_month, selected_competence_month, error = _resolve_competence_month(
        request.GET.get("competence_month")
    )
    if error:
        messages.error(request, error)
        return redirect(_timesheet_redirect_with_filters())

    is_month_closed = TimesheetMonthClosure.objects.filter(
        competence_month=competence_month
    ).exists()
    is_month_locked_by_bank_hours_closure = _is_month_locked_by_bank_hours_semester_closure(
        competence_month
    )
    pending_import_report = (
        TimesheetImportReport.objects.filter(
            competence_month=competence_month,
            status=TimesheetImportReport.STATUS_PENDING,
        )
        .order_by("-created_at", "-id")
        .first()
    )
    latest_import_report = (
        TimesheetImportReport.objects.filter(competence_month=competence_month)
        .order_by("-created_at", "-id")
        .first()
    )

    month_start, month_end = _get_month_date_range(competence_month)
    base_employees = _get_timesheet_base_employees()
    base_employee_ids = [employee.id for employee in base_employees]
    lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
    employees = _filter_employees_active_in_period(
        base_employees,
        month_start,
        month_end,
        lifecycle_dates,
    )
    employee_ids = [employee.id for employee in employees]
    schedule_events_by_employee = _build_employee_schedule_events_by_employee(employee_ids)
    month_entries = EmployeeTimeEntry.objects.select_related("employee").filter(
        competence_month=competence_month,
        employee_id__in=employee_ids,
    )
    entries_by_employee = {entry.employee_id: entry for entry in month_entries}

    calendar_ids = _collect_calendar_ids_for_timesheet(
        employees,
        schedule_events_by_employee,
    )

    calendar_exception_dates_by_calendar = _build_calendar_exception_dates_by_calendar(
        calendar_ids,
        month_start,
        month_end,
    )
    employee_exception_dates_by_employee = _build_employee_exception_dates_by_employee(
        employee_ids,
        month_start,
        month_end,
    )

    rows = []
    total_expected_minutes = 0
    for employee in employees:
        expected_minutes = _calculate_expected_minutes_for_employee(
            employee,
            month_start,
            month_end,
            calendar_exception_dates_by_calendar,
            employee_exception_dates_by_employee,
            schedule_events_by_employee,
            lifecycle_dates,
        )
        total_expected_minutes += expected_minutes
        rows.append(
            {
                "employee": employee,
                "entry": entries_by_employee.get(employee.id),
                "expected_minutes": expected_minutes,
            }
        )

    timesheet_entries_list = list(entries_by_employee.values())
    total_regular_minutes = sum(
        (entry.regular_minutes for entry in timesheet_entries_list),
        0,
    )
    total_overtime_60_minutes = sum(
        (entry.overtime_60_minutes for entry in timesheet_entries_list),
        0,
    )
    total_overtime_100_minutes = sum(
        (entry.overtime_100_minutes for entry in timesheet_entries_list),
        0,
    )
    total_overtime_minutes = total_overtime_60_minutes + total_overtime_100_minutes
    total_worked_minutes = total_regular_minutes + total_overtime_minutes
    total_absence_unexcused_minutes = sum(
        (entry.absence_unexcused_minutes for entry in timesheet_entries_list),
        0,
    )
    total_absence_excused_minutes = sum(
        (entry.absence_excused_minutes for entry in timesheet_entries_list),
        0,
    )
    total_absence_bank_minutes = sum(
        (
            entry.absence_bank_minutes
            if entry.employee.regime_compensacao_jornada
            == Employee.REGIME_COMPENSACAO_PARTICIPANTE
            else 0
            for entry in timesheet_entries_list
        ),
        0,
    )
    total_absence_minutes = (
        total_absence_unexcused_minutes
        + total_absence_excused_minutes
        + total_absence_bank_minutes
    )
    return render(
        request,
        "timesheet.html",
        {
            "rows": rows,
            "selected_competence_month": selected_competence_month,
            "is_month_closed": is_month_closed,
            "is_month_locked_by_bank_hours_closure": is_month_locked_by_bank_hours_closure,
            "registered_employees_count": len(entries_by_employee),
            "employees_count": len(employees),
            "total_regular_minutes": total_regular_minutes,
            "total_worked_minutes": total_worked_minutes,
            "total_overtime_minutes": total_overtime_minutes,
            "total_overtime_60_minutes": total_overtime_60_minutes,
            "total_overtime_100_minutes": total_overtime_100_minutes,
            "total_absence_unexcused_minutes": total_absence_unexcused_minutes,
            "total_absence_excused_minutes": total_absence_excused_minutes,
            "total_absence_bank_minutes": total_absence_bank_minutes,
            "total_absence_minutes": total_absence_minutes,
            "total_expected_minutes": total_expected_minutes,
            "pending_import_report": pending_import_report,
            "latest_import_report": latest_import_report,
        },
    )


def timesheet_import_report_page(request, report_id):
    report = get_object_or_404(TimesheetImportReport, id=report_id)
    competence_month_value = report.competence_month.strftime("%Y-%m")
    competence_month_label = _format_competence_month_label(report.competence_month)

    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()

        if action == "reject":
            with transaction.atomic():
                report = TimesheetImportReport.objects.select_for_update().get(id=report.id)
                if report.status != TimesheetImportReport.STATUS_PENDING:
                    messages.error(request, "Apenas relatorios pendentes podem ser rejeitados.")
                else:
                    report.status = TimesheetImportReport.STATUS_REJECTED
                    report.rejected_at = timezone.now()
                    report.save(update_fields=("status", "rejected_at"))
                    messages.success(request, "Relatorio rejeitado. Nenhum dado foi preenchido no timesheet.")
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        if action == "approve":
            with transaction.atomic():
                report = TimesheetImportReport.objects.select_for_update().get(id=report.id)
                if report.status != TimesheetImportReport.STATUS_PENDING:
                    messages.error(request, "Apenas relatorios pendentes podem ser aprovados.")
                    return redirect(
                        reverse(
                            "timesheet_import_report_page",
                            kwargs={"report_id": report.id},
                        )
                    )

                if TimesheetMonthClosure.objects.filter(
                    competence_month=report.competence_month
                ).exists():
                    messages.error(
                        request,
                        f"O timesheet de {competence_month_label} esta encerrado. Reabra o mes para aprovar a importacao.",
                    )
                    return redirect(
                        reverse(
                            "timesheet_import_report_page",
                            kwargs={"report_id": report.id},
                        )
                    )

                imported_count, apply_error = _apply_timesheet_import_report(report)
                if apply_error:
                    messages.error(request, apply_error)
                    return redirect(
                        reverse(
                            "timesheet_import_report_page",
                            kwargs={"report_id": report.id},
                        )
                    )

                report.status = TimesheetImportReport.STATUS_APPROVED
                report.approved_at = timezone.now()
                report.applied_entries_count = imported_count
                report.save(
                    update_fields=(
                        "status",
                        "approved_at",
                        "applied_entries_count",
                    )
                )
                messages.success(
                    request,
                    f"Relatorio aprovado. Importacao concluida para {imported_count} empregado(s).",
                )
            return redirect(
                _timesheet_redirect_with_filters(competence_month=competence_month_value)
            )

        messages.error(request, "Acao de relatorio invalida.")
        return redirect(
            reverse("timesheet_import_report_page", kwargs={"report_id": report.id})
        )

    warnings, hard_errors = _timesheet_import_report_messages(report)
    report_rows = list(report.rows.select_related("employee").all())
    ready_rows = [
        row for row in report_rows if row.row_type == TimesheetImportReportRow.TYPE_IMPORTED
    ]
    problem_rows = [
        row for row in report_rows if row.row_type == TimesheetImportReportRow.TYPE_SKIPPED
    ]
    missing_rows = [
        row
        for row in report_rows
        if row.row_type == TimesheetImportReportRow.TYPE_MISSING_IN_IMPORT
    ]
    is_month_closed = TimesheetMonthClosure.objects.filter(
        competence_month=report.competence_month
    ).exists()
    can_approve = (
        report.status == TimesheetImportReport.STATUS_PENDING
        and bool(ready_rows)
        and not hard_errors
        and not is_month_closed
    )

    return render(
        request,
        "timesheet_import_report.html",
        {
            "report": report,
            "selected_competence_month": competence_month_value,
            "competence_month_label": competence_month_label,
            "warnings": warnings,
            "hard_errors": hard_errors,
            "ready_rows": _decorate_timesheet_import_report_rows(ready_rows),
            "problem_rows": _decorate_timesheet_import_report_rows(problem_rows),
            "missing_rows": _decorate_timesheet_import_report_rows(missing_rows),
            "ready_rows_count": len(ready_rows),
            "problem_rows_count": len(problem_rows),
            "missing_rows_count": len(missing_rows),
            "is_month_closed": is_month_closed,
            "can_approve": can_approve,
        },
    )


def timesheet_snapshots_audit_page(request):
    competence_month, selected_competence_month, error = _resolve_competence_month(
        request.GET.get("competence_month")
    )
    if error:
        messages.error(request, error)
        return redirect(_snapshots_audit_redirect_with_filters())

    closure = TimesheetMonthClosure.objects.filter(competence_month=competence_month).first()

    schedule_snapshots = []
    calendar_period_snapshots = []
    employee_event_snapshots = []
    if closure:
        schedule_snapshots = list(
            TimesheetMonthClosureWorkScheduleSnapshot.objects.filter(closure=closure).order_by(
                "work_schedule_name",
                "source_work_schedule_id",
            )
        )
        calendar_period_snapshots = list(
            TimesheetMonthClosureCalendarPeriodSnapshot.objects.filter(closure=closure).order_by(
                "calendar_name",
                "start_date",
                "id",
            )
        )
        employee_event_snapshots = list(
            TimesheetMonthClosureEmployeeEventSnapshot.objects.filter(closure=closure).order_by(
                "employee_name",
                "effective_date",
                "source_employee_event_id",
            )
        )

    closed_months = list(
        TimesheetMonthClosure.objects.order_by("-competence_month").values_list(
            "competence_month", flat=True
        )[:24]
    )

    return render(
        request,
        "timesheet_snapshots_audit.html",
        {
            "selected_competence_month": selected_competence_month,
            "closure": closure,
            "schedule_snapshots": schedule_snapshots,
            "calendar_period_snapshots": calendar_period_snapshots,
            "employee_event_snapshots": employee_event_snapshots,
            "closed_months": closed_months,
            "closed_months_count": len(closed_months),
        },
    )


def timesheet_dashboard_page(request):
    start_month, end_month, selected_start_month, selected_end_month, error = (
        _resolve_competence_month_interval(
            request.GET.get("start_month"),
            request.GET.get("end_month"),
            request.GET.get("competence_month"),
        )
    )
    if error:
        messages.error(request, error)
        return redirect(_timesheet_redirect_with_filters())

    range_end = _get_month_date_range(end_month)[1]
    base_employees = _get_timesheet_base_employees()
    base_employee_ids = [employee.id for employee in base_employees]
    lifecycle_dates = _build_employee_lifecycle_dates(base_employee_ids)
    employees = _filter_employees_active_in_period(
        base_employees,
        start_month,
        range_end,
        lifecycle_dates,
    )
    employee_ids = [employee.id for employee in employees]
    schedule_events_by_employee = _build_employee_schedule_events_by_employee(employee_ids)

    competence_months = list(_iterate_competence_months(start_month, end_month))
    month_entries = EmployeeTimeEntry.objects.select_related("employee", "employee__sector").filter(
        competence_month__in=competence_months,
        employee_id__in=employee_ids,
    )
    entries_by_employee = {}
    for entry in month_entries:
        employee_entry = entries_by_employee.setdefault(
            entry.employee_id,
            {
                "regular_minutes": 0,
                "overtime_60_minutes": 0,
                "overtime_100_minutes": 0,
                "absence_unexcused_minutes": 0,
                "absence_excused_minutes": 0,
                "absence_bank_minutes": 0,
            },
        )
        employee_entry["regular_minutes"] += entry.regular_minutes
        employee_entry["overtime_60_minutes"] += entry.overtime_60_minutes
        employee_entry["overtime_100_minutes"] += entry.overtime_100_minutes
        employee_entry["absence_unexcused_minutes"] += entry.absence_unexcused_minutes
        employee_entry["absence_excused_minutes"] += entry.absence_excused_minutes
        employee_entry["absence_bank_minutes"] += entry.absence_bank_minutes

    calendar_ids = _collect_calendar_ids_for_timesheet(
        employees,
        schedule_events_by_employee,
    )

    calendar_exception_dates_by_calendar = _build_calendar_exception_dates_by_calendar(
        calendar_ids,
        start_month,
        range_end,
    )
    employee_exception_dates_by_employee = _build_employee_exception_dates_by_employee(
        employee_ids,
        start_month,
        range_end,
    )

    def _empty_group_bucket(name):
        return {
            "name": name,
            "expected_minutes": 0,
            "regular_minutes": 0,
            "overtime_60_minutes": 0,
            "overtime_100_minutes": 0,
            "worked_minutes": 0,
            "absence_unexcused_minutes": 0,
            "absence_excused_minutes": 0,
            "absence_bank_minutes": 0,
        }

    def _add_to_group_bucket(bucket, expected_minutes, regular_minutes, overtime_60_minutes, overtime_100_minutes, worked_minutes, absence_unexcused_minutes, absence_excused_minutes, absence_bank_minutes):
        bucket["expected_minutes"] += expected_minutes
        bucket["regular_minutes"] += regular_minutes
        bucket["overtime_60_minutes"] += overtime_60_minutes
        bucket["overtime_100_minutes"] += overtime_100_minutes
        bucket["worked_minutes"] += worked_minutes
        bucket["absence_unexcused_minutes"] += absence_unexcused_minutes
        bucket["absence_excused_minutes"] += absence_excused_minutes
        bucket["absence_bank_minutes"] += absence_bank_minutes

    def _build_group_rows(group_data, limit=None):
        rows = sorted(
            group_data.values(),
            key=lambda row: row["expected_minutes"],
            reverse=True,
        )
        if limit:
            rows = rows[:limit]
        for row in rows:
            row["balance_minutes"] = row["worked_minutes"] - row["expected_minutes"]
            row["expected_hhmm"] = _format_minutes_as_hour_label(row["expected_minutes"])
            row["worked_hhmm"] = _format_minutes_as_hour_label(row["worked_minutes"])
            row["overtime_60_hhmm"] = _format_minutes_as_hour_label(row["overtime_60_minutes"])
            row["overtime_100_hhmm"] = _format_minutes_as_hour_label(row["overtime_100_minutes"])
            row["balance_hhmm"] = _format_signed_minutes_as_hour_label(row["balance_minutes"])
            row["absence_unexcused_hhmm"] = _format_minutes_as_hour_label(
                row["absence_unexcused_minutes"]
            )
            row["absence_excused_hhmm"] = _format_minutes_as_hour_label(
                row["absence_excused_minutes"]
            )
            row["absence_bank_hhmm"] = _format_minutes_as_hour_label(
                row["absence_bank_minutes"]
            )
            row["utilization_percent"] = (
                round((row["worked_minutes"] * 100) / row["expected_minutes"])
                if row["expected_minutes"] > 0
                else 0
            )
        return rows

    sector_data = {}
    department_data = {}
    area_data = {}
    total_expected_minutes = 0
    total_regular_minutes = 0
    total_overtime_60_minutes = 0
    total_overtime_100_minutes = 0
    total_absence_unexcused_minutes = 0
    total_absence_excused_minutes = 0
    total_absence_bank_minutes = 0
    total_mod_minutes = 0
    total_moi_minutes = 0

    for employee in employees:
        expected_minutes = 0
        for competence_month in competence_months:
            month_start, month_end = _get_month_date_range(competence_month)
            expected_minutes += _calculate_expected_minutes_for_employee(
                employee,
                month_start,
                month_end,
                calendar_exception_dates_by_calendar,
                employee_exception_dates_by_employee,
                schedule_events_by_employee,
                lifecycle_dates,
            )
        total_expected_minutes += expected_minutes

        entry = entries_by_employee.get(employee.id)
        regular_minutes = entry["regular_minutes"] if entry else 0
        overtime_60_minutes = entry["overtime_60_minutes"] if entry else 0
        overtime_100_minutes = entry["overtime_100_minutes"] if entry else 0
        absence_unexcused_minutes = entry["absence_unexcused_minutes"] if entry else 0
        absence_excused_minutes = entry["absence_excused_minutes"] if entry else 0
        absence_bank_minutes = 0
        if entry and employee.regime_compensacao_jornada == Employee.REGIME_COMPENSACAO_PARTICIPANTE:
            absence_bank_minutes = entry["absence_bank_minutes"]

        worked_minutes = regular_minutes + overtime_60_minutes + overtime_100_minutes

        total_regular_minutes += regular_minutes
        total_overtime_60_minutes += overtime_60_minutes
        total_overtime_100_minutes += overtime_100_minutes
        total_absence_unexcused_minutes += absence_unexcused_minutes
        total_absence_excused_minutes += absence_excused_minutes
        total_absence_bank_minutes += absence_bank_minutes

        if employee.tipo == Employee.TYPE_DIRETO:
            total_mod_minutes += worked_minutes
        else:
            total_moi_minutes += worked_minutes

        if employee.sector and employee.sector.department:
            area_part = (
                employee.sector.department.area.nome
                if employee.sector.department.area
                else "Sem area"
            )
            department_part = employee.sector.department.nome
            sector_part = employee.sector.nome
            sector_key = f"{area_part}::{department_part}::{sector_part}"
            sector_name = f"{department_part} / {sector_part}"
        else:
            sector_key = "sem_departamento::sem_setor"
            sector_name = "Sem departamento / Sem setor"

        sector_bucket = sector_data.setdefault(sector_key, _empty_group_bucket(sector_name))
        _add_to_group_bucket(
            sector_bucket,
            expected_minutes,
            regular_minutes,
            overtime_60_minutes,
            overtime_100_minutes,
            worked_minutes,
            absence_unexcused_minutes,
            absence_excused_minutes,
            absence_bank_minutes,
        )

        department_name = "Sem departamento"
        area_name = "Sem area"
        if employee.sector and employee.sector.department:
            department_name = employee.sector.department.nome
            if employee.sector.department.area:
                area_name = employee.sector.department.area.nome

        department_bucket = department_data.setdefault(
            department_name, _empty_group_bucket(department_name)
        )
        _add_to_group_bucket(
            department_bucket,
            expected_minutes,
            regular_minutes,
            overtime_60_minutes,
            overtime_100_minutes,
            worked_minutes,
            absence_unexcused_minutes,
            absence_excused_minutes,
            absence_bank_minutes,
        )

        area_bucket = area_data.setdefault(area_name, _empty_group_bucket(area_name))
        _add_to_group_bucket(
            area_bucket,
            expected_minutes,
            regular_minutes,
            overtime_60_minutes,
            overtime_100_minutes,
            worked_minutes,
            absence_unexcused_minutes,
            absence_excused_minutes,
            absence_bank_minutes,
        )

    total_overtime_minutes = total_overtime_60_minutes + total_overtime_100_minutes
    total_worked_minutes = total_regular_minutes + total_overtime_minutes
    total_absence_minutes = (
        total_absence_unexcused_minutes
        + total_absence_excused_minutes
        + total_absence_bank_minutes
    )

    total_type_minutes = total_mod_minutes + total_moi_minutes
    mod_share_percent = (
        round((total_mod_minutes * 100) / total_type_minutes) if total_type_minutes else 0
    )
    moi_share_percent = 100 - mod_share_percent if total_type_minutes else 0

    sector_rows = _build_group_rows(sector_data)
    sector_rows = sorted(sector_rows, key=lambda row: row["name"].casefold())[:12]
    department_rows = _build_group_rows(department_data)
    area_rows = _build_group_rows(area_data)

    top_expected_minutes = max((row["expected_minutes"] for row in sector_rows), default=0)

    chart_labels = [row["name"] for row in sector_rows]
    chart_expected_hours = [round(row["expected_minutes"] / 60, 2) for row in sector_rows]
    chart_worked_hours = [round(row["regular_minutes"] / 60, 2) for row in sector_rows]
    chart_overtime_60_hours = [round(row["overtime_60_minutes"] / 60, 2) for row in sector_rows]
    chart_overtime_100_hours = [round(row["overtime_100_minutes"] / 60, 2) for row in sector_rows]
    chart_balance_hours = [
        round((row["worked_minutes"] - row["expected_minutes"]) / 60, 2)
        for row in sector_rows
    ]
    chart_absence_hours = [
        round(row["absence_unexcused_minutes"] / 60, 2) for row in sector_rows
    ]
    chart_excused_hours = [
        round(row["absence_excused_minutes"] / 60, 2) for row in sector_rows
    ]
    department_chart_rows = department_rows[:12]
    area_chart_rows = area_rows[:12]
    department_chart_labels = [row["name"] for row in department_chart_rows]
    department_chart_expected_hours = [
        round(row["expected_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_worked_hours = [
        round(row["regular_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_overtime_60_hours = [
        round(row["overtime_60_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_overtime_100_hours = [
        round(row["overtime_100_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_absence_unexcused_hours = [
        round(row["absence_unexcused_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_absence_excused_hours = [
        round(row["absence_excused_minutes"] / 60, 2) for row in department_chart_rows
    ]
    department_chart_absence_bank_hours = [
        round(row["absence_bank_minutes"] / 60, 2) for row in department_chart_rows
    ]
    area_chart_labels = [row["name"] for row in area_chart_rows]
    area_chart_expected_hours = [round(row["expected_minutes"] / 60, 2) for row in area_chart_rows]
    area_chart_worked_hours = [round(row["regular_minutes"] / 60, 2) for row in area_chart_rows]
    area_chart_overtime_60_hours = [
        round(row["overtime_60_minutes"] / 60, 2) for row in area_chart_rows
    ]
    area_chart_overtime_100_hours = [
        round(row["overtime_100_minutes"] / 60, 2) for row in area_chart_rows
    ]

    range_start = _get_month_date_range(start_month)[0]
    range_end = _get_month_date_range(end_month)[1]
    month_reference_label = f"{range_start.strftime('%d/%m')} - {range_end.strftime('%d/%m/%Y')}"
    if start_month == end_month:
        competence_month_label = _format_competence_month_label(start_month)
    else:
        competence_month_label = (
            f"{_format_competence_month_label(start_month)} a "
            f"{_format_competence_month_label(end_month)}"
        )

    return render(
        request,
        "dashboard_horas_grouped_stacked.html",
        {
            "selected_competence_month": selected_end_month,
            "selected_start_month": selected_start_month,
            "selected_end_month": selected_end_month,
            "competence_month_label": competence_month_label,
            "month_reference_label": month_reference_label,
            "employees_count": len(employees),
            "registered_employees_count": len(entries_by_employee),
            "total_expected_minutes": total_expected_minutes,
            "total_regular_minutes": total_regular_minutes,
            "total_worked_minutes": total_worked_minutes,
            "total_overtime_minutes": total_overtime_minutes,
            "total_absence_minutes": total_absence_minutes,
            "total_absence_unexcused_minutes": total_absence_unexcused_minutes,
            "total_absence_excused_minutes": total_absence_excused_minutes,
            "total_absence_bank_minutes": total_absence_bank_minutes,
            "mod_share_percent": mod_share_percent,
            "moi_share_percent": moi_share_percent,
            "sector_rows": sector_rows,
            "department_rows": department_rows,
            "area_rows": area_rows,
            "top_expected_minutes": top_expected_minutes,
            "chart_labels": chart_labels,
            "chart_expected_hours": chart_expected_hours,
            "chart_worked_hours": chart_worked_hours,
            "chart_overtime_60_hours": chart_overtime_60_hours,
            "chart_overtime_100_hours": chart_overtime_100_hours,
            "chart_balance_hours": chart_balance_hours,
            "chart_absence_hours": chart_absence_hours,
            "chart_excused_hours": chart_excused_hours,
            "department_chart_labels": department_chart_labels,
            "department_chart_expected_hours": department_chart_expected_hours,
            "department_chart_worked_hours": department_chart_worked_hours,
            "department_chart_overtime_60_hours": department_chart_overtime_60_hours,
            "department_chart_overtime_100_hours": department_chart_overtime_100_hours,
            "department_chart_absence_unexcused_hours": department_chart_absence_unexcused_hours,
            "department_chart_absence_excused_hours": department_chart_absence_excused_hours,
            "department_chart_absence_bank_hours": department_chart_absence_bank_hours,
            "area_chart_labels": area_chart_labels,
            "area_chart_expected_hours": area_chart_expected_hours,
            "area_chart_worked_hours": area_chart_worked_hours,
            "area_chart_overtime_60_hours": area_chart_overtime_60_hours,
            "area_chart_overtime_100_hours": area_chart_overtime_100_hours,
            "total_expected_hhmm": _format_minutes_as_hour_label(total_expected_minutes),
            "total_worked_hhmm": _format_minutes_as_hour_label(total_worked_minutes),
            "total_overtime_hhmm": _format_minutes_as_hour_label(total_overtime_minutes),
            "total_absence_hhmm": _format_minutes_as_hour_label(total_absence_minutes),
        },
    )


def bank_hours_page(request):
    start_month, end_month, selected_start_month, selected_end_month, error = (
        _resolve_competence_month_interval(
            request.GET.get("start_month"),
            request.GET.get("end_month"),
            request.GET.get("competence_month"),
        )
    )
    if error:
        messages.error(request, error)
        return redirect(_bank_hours_redirect_with_filters())

    participants = list(
        Employee.objects.select_related("sector")
        .filter(
            deactivated_at__isnull=True,
            regime_compensacao_jornada=Employee.REGIME_COMPENSACAO_PARTICIPANTE,
        )
        .order_by("nome_completo")
    )

    competence_months = list(_iterate_competence_months(start_month, end_month))
    competence_month_set = set(competence_months)
    participant_ids = [employee.id for employee in participants]

    entries_qs = EmployeeTimeEntry.objects.filter(employee_id__in=participant_ids)
    entries_in_range = entries_qs.filter(competence_month__in=competence_months)
    entries_before_range = entries_qs.filter(competence_month__lt=start_month)

    monthly_entry_data = defaultdict(
        lambda: {
            "overtime_minutes": 0,
            "absence_minutes": 0,
            "manual_minutes": 0,
        }
    )
    opening_balance_by_employee = defaultdict(int)
    opening_balance_known_by_employee = defaultdict(lambda: True)

    for entry in entries_in_range:
        key = (entry.employee_id, entry.competence_month)
        overtime_minutes = entry.overtime_60_minutes + entry.overtime_100_minutes
        absence_minutes = (
            entry.absence_unexcused_minutes
            + entry.absence_bank_minutes
        )
        monthly_entry_data[key]["overtime_minutes"] += overtime_minutes
        monthly_entry_data[key]["absence_minutes"] += absence_minutes

    for entry in entries_before_range:
        competence_month = entry.competence_month
        overtime_minutes = entry.overtime_60_minutes + entry.overtime_100_minutes
        absence_minutes = (
            entry.absence_unexcused_minutes
            + entry.absence_bank_minutes
        )
        formulas = _resolve_bank_hours_formulas_for_month(competence_month)
        calc = _calculate_bank_hours_columns(
            overtime_minutes=overtime_minutes,
            absence_minutes=absence_minutes,
            manual_minutes=0,
            formulas=formulas,
        )
        if calc["is_f_defined"]:
            opening_balance_by_employee[entry.employee_id] += calc["month_delta_minutes"]
        else:
            opening_balance_known_by_employee[entry.employee_id] = False

    movement_events_qs = EmployeeEvent.objects.filter(
        employee_id__in=participant_ids,
        event_type=EmployeeEvent.EVENT_TYPE_BANK_HOURS_MOVEMENT,
        bank_hours_amount__isnull=False,
    )

    for movement in movement_events_qs.filter(
        effective_date__lt=start_month,
    ):
        opening_balance_by_employee[movement.employee_id] += _minutes_from_decimal_hours(
            movement.bank_hours_amount
        )

    for movement in movement_events_qs.filter(
        effective_date__gte=start_month,
        effective_date__lte=_get_month_date_range(end_month)[1],
    ):
        month_key = _month_start_for_date(movement.effective_date)
        if month_key not in competence_month_set:
            continue
        monthly_entry_data[(movement.employee_id, month_key)]["manual_minutes"] += (
            _minutes_from_decimal_hours(movement.bank_hours_amount)
        )

    rows = []
    total_credit_minutes = 0
    total_hour_bank_base_minutes = 0
    total_hour_bank_bonus_minutes = 0
    total_absence_minutes = 0
    total_manual_minutes = 0
    total_month_delta_minutes = 0
    total_closing_balance_minutes = 0
    has_missing_rule_b = False
    has_missing_rule_c = False
    has_missing_rule_f = False
    applied_formulas_map = {}
    employee_condensed_rows = []

    for employee in participants:
        running_balance = opening_balance_by_employee[employee.id]
        running_balance_known = opening_balance_known_by_employee[employee.id]
        employee_total_hour_bank_base_minutes = 0
        employee_total_hour_bank_bonus_minutes = 0
        employee_total_absence_minutes = 0
        employee_total_manual_minutes = 0
        employee_total_month_delta_minutes = 0
        employee_missing_rule_b = False
        employee_missing_rule_c = False
        employee_missing_rule_f = False
        for competence_month in competence_months:
            payload = monthly_entry_data[(employee.id, competence_month)]
            formulas = _resolve_bank_hours_formulas_for_month(competence_month)
            applied_formulas_map[competence_month] = formulas
            calc = _calculate_bank_hours_columns(
                overtime_minutes=payload["overtime_minutes"],
                absence_minutes=payload["absence_minutes"],
                manual_minutes=payload["manual_minutes"],
                formulas=formulas,
            )
            month_delta_minutes = calc["month_delta_minutes"] if calc["is_f_defined"] else None
            if calc["is_f_defined"] and running_balance_known:
                running_balance += month_delta_minutes
            else:
                running_balance_known = False

            has_missing_rule_b = has_missing_rule_b or (not calc["is_b_defined"])
            has_missing_rule_c = has_missing_rule_c or (not calc["is_c_defined"])
            has_missing_rule_f = has_missing_rule_f or (not calc["is_f_defined"])
            employee_missing_rule_b = employee_missing_rule_b or (not calc["is_b_defined"])
            employee_missing_rule_c = employee_missing_rule_c or (not calc["is_c_defined"])
            employee_missing_rule_f = employee_missing_rule_f or (not calc["is_f_defined"])
            rows.append(
                {
                    "employee": employee,
                    "competence_month": competence_month,
                    "overtime_minutes": calc["overtime_minutes"],
                    "overtime_hhmm": _format_minutes_as_hour_label(
                        calc["overtime_minutes"]
                    ),
                    "hour_bank_base_minutes": calc["hour_bank_base_minutes"],
                    "hour_bank_base_hhmm": _format_minutes_as_hour_label(
                        calc["hour_bank_base_minutes"]
                    ),
                    "is_b_defined": calc["is_b_defined"],
                    "hour_bank_bonus_minutes": calc["hour_bank_bonus_minutes"],
                    "hour_bank_bonus_hhmm": _format_minutes_as_hour_label(
                        calc["hour_bank_bonus_minutes"]
                    ),
                    "is_c_defined": calc["is_c_defined"],
                    "credit_minutes": calc["credit_minutes"],
                    "credit_hhmm": _format_minutes_as_hour_label(
                        calc["credit_minutes"]
                    ),
                    "absence_minutes": calc["absence_minutes"],
                    "absence_hhmm": _format_minutes_as_hour_label(
                        calc["absence_minutes"]
                    ),
                    "manual_minutes": calc["manual_minutes"],
                    "manual_hhmm": _format_signed_minutes_as_hour_label(
                        calc["manual_minutes"]
                    ),
                    "month_delta_minutes": month_delta_minutes if month_delta_minutes is not None else 0,
                    "month_delta_hhmm": _format_signed_minutes_as_hour_label(month_delta_minutes or 0),
                    "is_f_defined": calc["is_f_defined"],
                    "running_balance_minutes": running_balance if running_balance_known else 0,
                    "running_balance_hhmm": _format_signed_minutes_as_hour_label(running_balance if running_balance_known else 0),
                    "running_balance_defined": running_balance_known,
                }
            )
            total_credit_minutes += calc["credit_minutes"]
            total_hour_bank_base_minutes += calc["hour_bank_base_minutes"]
            total_hour_bank_bonus_minutes += calc["hour_bank_bonus_minutes"]
            total_absence_minutes += calc["absence_minutes"]
            total_manual_minutes += calc["manual_minutes"]
            employee_total_hour_bank_base_minutes += calc["hour_bank_base_minutes"]
            employee_total_hour_bank_bonus_minutes += calc["hour_bank_bonus_minutes"]
            employee_total_absence_minutes += calc["absence_minutes"]
            employee_total_manual_minutes += calc["manual_minutes"]
            if month_delta_minutes is not None:
                total_month_delta_minutes += month_delta_minutes
                employee_total_month_delta_minutes += month_delta_minutes

        if running_balance_known:
            total_closing_balance_minutes += running_balance

        employee_condensed_rows.append(
            {
                "employee": employee,
                "total_hour_bank_base_minutes": employee_total_hour_bank_base_minutes,
                "total_hour_bank_bonus_minutes": employee_total_hour_bank_bonus_minutes,
                "total_absence_minutes": employee_total_absence_minutes,
                "total_manual_minutes": employee_total_manual_minutes,
                "total_month_delta_minutes": employee_total_month_delta_minutes,
                "closing_balance_minutes": running_balance if running_balance_known else 0,
                "closing_balance_defined": running_balance_known,
                "has_missing_rule_b": employee_missing_rule_b,
                "has_missing_rule_c": employee_missing_rule_c,
                "has_missing_rule_f": employee_missing_rule_f,
            }
        )

    return render(
        request,
        "bank_hours.html",
        {
            "rows": rows,
            "participants": participants,
            "participants_count": len(participants),
            "selected_start_month": selected_start_month,
            "selected_end_month": selected_end_month,
            "total_credit_minutes": total_credit_minutes,
            "total_credit_hhmm": _format_minutes_as_hour_label(total_credit_minutes),
            "total_hour_bank_base_minutes": total_hour_bank_base_minutes,
            "total_hour_bank_base_hhmm": _format_minutes_as_hour_label(
                total_hour_bank_base_minutes
            ),
            "total_hour_bank_bonus_minutes": total_hour_bank_bonus_minutes,
            "total_hour_bank_bonus_hhmm": _format_minutes_as_hour_label(
                total_hour_bank_bonus_minutes
            ),
            "total_absence_minutes": total_absence_minutes,
            "total_absence_hhmm": _format_minutes_as_hour_label(total_absence_minutes),
            "total_manual_minutes": total_manual_minutes,
            "total_manual_hhmm": _format_signed_minutes_as_hour_label(
                total_manual_minutes
            ),
            "total_month_delta_minutes": total_month_delta_minutes,
            "total_month_delta_hhmm": _format_signed_minutes_as_hour_label(
                total_month_delta_minutes
            ),
            "total_closing_balance_minutes": total_closing_balance_minutes,
            "total_closing_balance_hhmm": _format_signed_minutes_as_hour_label(
                total_closing_balance_minutes
            ),
            "rule_example": "Sem regra padrao. Cadastre formulas para B, C e F.",
            "has_missing_rule_b": has_missing_rule_b,
            "has_missing_rule_c": has_missing_rule_c,
            "has_missing_rule_f": has_missing_rule_f,
            "applied_formulas_map": applied_formulas_map,
            "employee_condensed_rows": employee_condensed_rows,
        },
    )


def bank_hours_closure_page(request):
    today = timezone.localdate()
    raw_year = (request.POST.get("year") or request.GET.get("year") or "").strip()
    raw_semester = (request.POST.get("semester") or request.GET.get("semester") or "").strip()

    selected_year = today.year
    if raw_year.isdigit():
        parsed_year = int(raw_year)
        if 2000 <= parsed_year <= 2100:
            selected_year = parsed_year

    if raw_semester in {"1", "2"}:
        selected_semester = int(raw_semester)
    else:
        selected_semester = 1 if today.month <= 6 else 2

    semester_months, competence_months, start_month, end_month, adjustment_date = _resolve_semester_window(
        selected_year,
        selected_semester,
    )
    closed_months = set(
        TimesheetMonthClosure.objects.filter(competence_month__in=competence_months).values_list(
            "competence_month",
            flat=True,
        )
    )
    missing_closures = [
        competence_month
        for competence_month in competence_months
        if competence_month not in closed_months
    ]

    participants = list(
        Employee.objects.select_related("sector")
        .filter(regime_compensacao_jornada=Employee.REGIME_COMPENSACAO_PARTICIPANTE)
        .order_by("nome_completo")
    )
    participant_ids = [employee.id for employee in participants]
    competence_month_set = set(competence_months)

    entries_qs = EmployeeTimeEntry.objects.filter(employee_id__in=participant_ids)
    entries_in_range = entries_qs.filter(competence_month__in=competence_months)
    entries_before_range = entries_qs.filter(competence_month__lt=start_month)

    monthly_entry_data = defaultdict(
        lambda: {
            "overtime_minutes": 0,
            "absence_minutes": 0,
            "manual_minutes": 0,
        }
    )
    opening_balance_by_employee = defaultdict(int)
    opening_balance_known_by_employee = defaultdict(lambda: True)

    for entry in entries_in_range:
        key = (entry.employee_id, entry.competence_month)
        overtime_minutes = entry.overtime_60_minutes + entry.overtime_100_minutes
        absence_minutes = entry.absence_unexcused_minutes + entry.absence_bank_minutes
        monthly_entry_data[key]["overtime_minutes"] += overtime_minutes
        monthly_entry_data[key]["absence_minutes"] += absence_minutes

    for entry in entries_before_range:
        formulas = _resolve_bank_hours_formulas_for_month(entry.competence_month)
        calc = _calculate_bank_hours_columns(
            overtime_minutes=entry.overtime_60_minutes + entry.overtime_100_minutes,
            absence_minutes=entry.absence_unexcused_minutes + entry.absence_bank_minutes,
            manual_minutes=0,
            formulas=formulas,
        )
        if calc["is_f_defined"]:
            opening_balance_by_employee[entry.employee_id] += calc["month_delta_minutes"]
        else:
            opening_balance_known_by_employee[entry.employee_id] = False

    movement_events_qs = EmployeeEvent.objects.filter(
        employee_id__in=participant_ids,
        event_type=EmployeeEvent.EVENT_TYPE_BANK_HOURS_MOVEMENT,
        bank_hours_amount__isnull=False,
    )
    for movement in movement_events_qs.filter(effective_date__lt=start_month):
        opening_balance_by_employee[movement.employee_id] += _minutes_from_decimal_hours(
            movement.bank_hours_amount
        )
    for movement in movement_events_qs.filter(
        effective_date__gte=start_month,
        effective_date__lte=_get_month_date_range(end_month)[1],
    ):
        month_key = _month_start_for_date(movement.effective_date)
        if month_key in competence_month_set:
            monthly_entry_data[(movement.employee_id, month_key)]["manual_minutes"] += _minutes_from_decimal_hours(
                movement.bank_hours_amount
            )

    closure_rows = []
    total_positive_minutes = 0
    total_negative_minutes = 0
    total_with_balance = 0
    unknown_balance_count = 0

    for employee in participants:
        running_balance = opening_balance_by_employee[employee.id]
        running_balance_known = opening_balance_known_by_employee[employee.id]
        for competence_month in competence_months:
            payload = monthly_entry_data[(employee.id, competence_month)]
            formulas = _resolve_bank_hours_formulas_for_month(competence_month)
            calc = _calculate_bank_hours_columns(
                overtime_minutes=payload["overtime_minutes"],
                absence_minutes=payload["absence_minutes"],
                manual_minutes=payload["manual_minutes"],
                formulas=formulas,
            )
            if calc["is_f_defined"] and running_balance_known:
                running_balance += calc["month_delta_minutes"]
            else:
                running_balance_known = False

        if running_balance_known:
            positive_minutes = running_balance if running_balance > 0 else 0
            negative_minutes = -running_balance if running_balance < 0 else 0
            total_positive_minutes += positive_minutes
            total_negative_minutes += negative_minutes
            total_with_balance += 1
        else:
            positive_minutes = 0
            negative_minutes = 0
            unknown_balance_count += 1

        closure_rows.append(
            {
                "employee": employee,
                "status_label": "Desligado" if employee.deactivated_at else "Ativo",
                "running_balance_defined": running_balance_known,
                "running_balance_minutes": running_balance if running_balance_known else 0,
                "running_balance_hhmm": (
                    _format_signed_minutes_as_hour_label(running_balance)
                    if running_balance_known
                    else "-"
                ),
                "positive_hhmm": (
                    _format_minutes_as_hour_label(positive_minutes)
                    if running_balance_known
                    else "-"
                ),
                "negative_hhmm": (
                _format_minutes_as_hour_label(negative_minutes)
                    if running_balance_known
                    else "-"
                ),
                "closing_after_hhmm": "00:00" if running_balance_known else "-",
            }
        )

    active_closure = (
        BankHoursSemesterClosure.objects.filter(
            year=selected_year,
            semester=selected_semester,
            reversed_at__isnull=True,
        )
        .order_by("-closed_at", "-id")
        .first()
    )

    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()

        if action == "close_semester":
            if missing_closures:
                messages.error(
                    request,
                    "Encerramento bloqueado: todos os meses do semestre precisam estar encerrados no timesheet.",
                )
            elif active_closure:
                messages.error(
                    request,
                    "Ja existe um encerramento ativo para este semestre. Estorne antes de encerrar novamente.",
                )
            else:
                adjustable_rows = [
                    row
                    for row in closure_rows
                    if row["running_balance_defined"] and row["running_balance_minutes"] != 0
                ]
                with transaction.atomic():
                    closure = BankHoursSemesterClosure.objects.create(
                        year=selected_year,
                        semester=selected_semester,
                        adjustment_month=end_month,
                    )
                    adjustments = []
                    for row in adjustable_rows:
                        employee = row["employee"]
                        balance_minutes = row["running_balance_minutes"]
                        adjustment_minutes = -balance_minutes
                        adjustment_hours = _decimal_hours_from_minutes(adjustment_minutes)
                        closing_event = EmployeeEvent.objects.create(
                            employee=employee,
                            event_type=EmployeeEvent.EVENT_TYPE_BANK_HOURS_MOVEMENT,
                            effective_date=adjustment_date,
                            bank_hours_amount=adjustment_hours,
                            notes=(
                                f"Ajuste automatico de encerramento semestral BH "
                                f"{selected_year}/S{selected_semester} (saldo base: "
                                f"{_format_signed_minutes_as_hour_label(balance_minutes)})."
                            ),
                        )
                        adjustments.append(
                            BankHoursSemesterClosureAdjustment(
                                closure=closure,
                                employee=employee,
                                balance_minutes=balance_minutes,
                                closing_event=closing_event,
                            )
                        )
                    if adjustments:
                        BankHoursSemesterClosureAdjustment.objects.bulk_create(adjustments)
                messages.success(
                    request,
                    (
                        f"Semestre {selected_year}/S{selected_semester} encerrado. "
                        f"{len(adjustable_rows)} ajuste(s) gerado(s) no mes {end_month.strftime('%m/%Y')}."
                    ),
                )
            return redirect(
                f"{reverse('bank_hours_closure_page')}?year={selected_year}&semester={selected_semester}"
            )

        if action == "reverse_semester":
            if not active_closure:
                messages.error(request, "Nao existe encerramento ativo para estornar neste semestre.")
            else:
                with transaction.atomic():
                    adjustments = list(
                        active_closure.adjustments.select_related("employee").order_by("id")
                    )
                    reversed_count = 0
                    for adjustment in adjustments:
                        if adjustment.reversal_event_id:
                            continue
                        reversal_hours = _decimal_hours_from_minutes(adjustment.balance_minutes)
                        reversal_event = EmployeeEvent.objects.create(
                            employee=adjustment.employee,
                            event_type=EmployeeEvent.EVENT_TYPE_BANK_HOURS_MOVEMENT,
                            effective_date=today,
                            bank_hours_amount=reversal_hours,
                            notes=(
                                f"Estorno automatico do encerramento semestral BH "
                                f"{active_closure.year}/S{active_closure.semester}."
                            ),
                        )
                        adjustment.reversal_event = reversal_event
                        adjustment.save(update_fields=["reversal_event"])
                        reversed_count += 1

                    active_closure.reversed_at = timezone.now()
                    active_closure.save(update_fields=["reversed_at"])

                messages.success(
                    request,
                    (
                        f"Encerramento {selected_year}/S{selected_semester} estornado com sucesso. "
                        f"{reversed_count} ajuste(s) revertido(s)."
                    ),
                )
            return redirect(
                f"{reverse('bank_hours_closure_page')}?year={selected_year}&semester={selected_semester}"
            )

    return render(
        request,
        "bank_hours_closure.html",
        {
            "selected_year": selected_year,
            "selected_semester": str(selected_semester),
            "can_close_bank_hours": not missing_closures,
            "missing_closure_months": [
                _format_competence_month_label(month_date) for month_date in missing_closures
            ],
            "participants_count": len(participants),
            "total_with_balance": total_with_balance,
            "unknown_balance_count": unknown_balance_count,
            "total_positive_hhmm": _format_minutes_as_hour_label(total_positive_minutes),
            "total_negative_hhmm": _format_minutes_as_hour_label(total_negative_minutes),
            "closure_rows": closure_rows,
            "active_closure": active_closure,
            "last_closed_month_label": end_month.strftime("%m/%Y"),
        },
    )


def bank_hours_rules_page(request):
    start_month, end_month, selected_start_month, selected_end_month, error = (
        _resolve_competence_month_interval(
            request.GET.get("start_month"),
            request.GET.get("end_month"),
            request.GET.get("competence_month"),
        )
    )
    if error:
        messages.error(request, error)
        return redirect(_bank_hours_rules_redirect_with_filters())

    if request.method == "POST":
        target_column = (request.POST.get("target_column") or "").strip().upper()
        raw_formula = request.POST.get("formula")
        raw_start = (request.POST.get("start_month") or "").strip()
        raw_end = (request.POST.get("end_month") or "").strip()

        if target_column not in {"B", "C", "F"}:
            messages.error(request, "Selecione uma coluna de destino valida.")
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )

        start_month_input, _, start_error = _resolve_competence_month(raw_start)
        if start_error:
            messages.error(request, "Informe um mes inicial valido.")
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )

        end_month_input, _, end_error = _resolve_competence_month(raw_end)
        if end_error:
            messages.error(request, "Informe um mes final valido.")
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )
        if end_month_input < start_month_input:
            messages.error(request, "Mes final deve ser maior ou igual ao mes inicial.")
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )

        formula, formula_error = _normalize_rule_formula(target_column, raw_formula)
        if formula_error:
            messages.error(request, formula_error)
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )

        sample_values = {
            "A": Decimal(120),
            "B": Decimal(60),
            "C": Decimal(36),
            "D": Decimal(30),
            "E": Decimal(0),
            "F": Decimal(6),
        }
        try:
            _safe_decimal_formula_eval(formula, sample_values)
        except (ValueError, InvalidOperation):
            messages.error(request, "Formula invalida. Use somente A..F, numeros e + - * /.")
            return redirect(
                _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
            )

        BankHoursRule.objects.create(
            start_month=start_month_input,
            end_month=end_month_input,
            target_column=target_column,
            formula=formula,
        )
        messages.success(request, "Regra cadastrada com sucesso.")
        return redirect(
            _bank_hours_rules_redirect_with_filters(selected_start_month, selected_end_month)
        )

    filter_column = (request.GET.get("filter_column") or "").strip().upper()
    filter_query = (request.GET.get("q") or "").strip()
    raw_filter_start_month = (request.GET.get("filter_start_month") or "").strip()
    raw_filter_end_month = (request.GET.get("filter_end_month") or "").strip()

    filter_start_month = None
    filter_end_month = None

    if raw_filter_start_month:
        filter_start_month, _, filter_start_error = _resolve_competence_month(raw_filter_start_month)
        if filter_start_error:
            messages.error(request, "Filtro de mes inicial invalido.")
            filter_start_month = None
            raw_filter_start_month = ""
    if raw_filter_end_month:
        filter_end_month, _, filter_end_error = _resolve_competence_month(raw_filter_end_month)
        if filter_end_error:
            messages.error(request, "Filtro de mes final invalido.")
            filter_end_month = None
            raw_filter_end_month = ""

    rules = BankHoursRule.objects.all()
    if filter_column in {"B", "C", "F"}:
        rules = rules.filter(target_column=filter_column)
    if filter_query:
        rules = rules.filter(formula__icontains=filter_query)
    if filter_start_month:
        rules = rules.filter(end_month__gte=filter_start_month)
    if filter_end_month:
        rules = rules.filter(start_month__lte=filter_end_month)
    rules = rules.order_by("-start_month", "target_column", "-id")

    return render(
        request,
        "bank_hours_rules.html",
        {
            "rules": rules,
            "selected_start_month": selected_start_month,
            "selected_end_month": selected_end_month,
            "default_formulas": {"B": "S/Regra", "C": "S/Regra", "F": "S/Regra"},
            "filter_column": filter_column,
            "filter_query": filter_query,
            "filter_start_month": raw_filter_start_month,
            "filter_end_month": raw_filter_end_month,
        },
    )


@require_POST
def bank_hours_rule_delete(request, rule_id):
    start_month = (request.POST.get("start_month") or "").strip()
    end_month = (request.POST.get("end_month") or "").strip()
    rule = get_object_or_404(BankHoursRule, id=rule_id)
    rule.delete()
    messages.success(request, "Regra removida.")
    return redirect(_bank_hours_rules_redirect_with_filters(start_month, end_month))




@require_POST
def employee_create_sector(request):
    _, department, nome, taxonomy_error = _parse_sector_taxonomy_or_error(request.POST)
    if taxonomy_error:
        messages.error(request, taxonomy_error)
        return redirect(
            _employees_redirect_with_flags(
                open_employee_modal=1,
                open_sector_modal=1,
            )
        )

    if Sector.objects.filter(
        department=department,
        nome__iexact=nome,
        deactivated_at__isnull=True,
    ).exists():
        messages.error(request, "Ja existe um setor com esse nome neste departamento.")
        return redirect(
            _employees_redirect_with_flags(
                open_employee_modal=1,
                open_sector_modal=1,
            )
        )

    sector = Sector.objects.create(nome=nome, department=department)
    messages.success(request, "Setor cadastrado com sucesso.")
    return redirect(
        _employees_redirect_with_flags(
            open_employee_modal=1,
            selected_sector=sector.id,
        )
    )


def employee_edit_page(request, employee_id):
    employee = get_object_or_404(
        Employee.objects.select_related("cargo", "sector__department__area", "work_schedule"),
        id=employee_id,
    )
    active_cargos = Cargo.objects.filter(deactivated_at__isnull=True).order_by("nome")
    active_sectors = (
        Sector.objects.filter(deactivated_at__isnull=True)
        .select_related("department__area")
        .order_by("department__area__nome", "department__nome", "nome")
    )
    active_areas, active_departments = _active_taxonomy()
    work_schedules = WorkSchedule.objects.all().order_by("nome")
    current_cargo = employee.cargo
    current_sector = employee.sector
    allocation_events = EmployeeEvent.objects.filter(
        employee=employee,
        event_type=EmployeeEvent.EVENT_TYPE_ALLOCATION_CHANGE,
    ).select_related(
        "previous_sector__department__area",
        "new_sector__department__area",
        "previous_work_schedule",
        "new_work_schedule",
    )
    sector_history = [
        event
        for event in allocation_events
        if (
            event.previous_sector_id != event.new_sector_id
            or (event.notes or "").strip()
        )
    ]
    work_schedule_history = [
        event
        for event in allocation_events
        if event.previous_work_schedule_id != event.new_work_schedule_id
    ]
    calendar_exception_types = (
        EmployeeEvent.EVENT_TYPE_DAY_OFF,
        EmployeeEvent.EVENT_TYPE_VACATION,
        EmployeeEvent.EVENT_TYPE_ABSENCE,
        EmployeeEvent.EVENT_TYPE_MEDICAL_CERTIFICATE,
    )
    calendar_exception_events = EmployeeEvent.objects.filter(
        employee=employee,
        event_type__in=calendar_exception_types,
    ).order_by("-effective_date", "-created_at", "-id")
    selected_tab = (request.GET.get("tab") or "").strip()
    if selected_tab not in ("info", "cargo-setor", "escala", "calendar-exceptions"):
        selected_tab = "info"

    if request.method == "POST":
        form_type = (request.POST.get("form_type") or "").strip()
        if form_type == "calendar_exception":
            event_type = (request.POST.get("exception_type") or "").strip()
            if event_type not in calendar_exception_types:
                messages.error(request, "Selecione um tipo de excessao valido.")
                return redirect(
                    _employee_edit_redirect_with_tab(employee.id, "calendar-exceptions")
                )

            start_date, end_date, date_error = _parse_special_period_dates(request.POST)
            if date_error:
                messages.error(request, date_error)
                return redirect(
                    _employee_edit_redirect_with_tab(employee.id, "calendar-exceptions")
                )

            closed_month = _find_first_closed_competence_month(start_date, end_date)
            if closed_month:
                messages.error(
                    request,
                    "Nao e permitido registrar evento/ajuste com data em competencia encerrada "
                    f"({_format_competence_month_label(closed_month)}).",
                )
                return redirect(
                    _employee_edit_redirect_with_tab(employee.id, "calendar-exceptions")
                )

            EmployeeEvent.objects.create(
                employee=employee,
                event_type=event_type,
                effective_date=start_date,
                end_date=end_date,
            )
            messages.success(request, "Excessao de calendario registrada com sucesso.")
            return redirect(
                _employee_edit_redirect_with_tab(employee.id, "calendar-exceptions")
            )
        matricula = (request.POST.get("matricula") or "").strip()
        nome_completo = (request.POST.get("nome_completo") or "").strip()
        tipo, regime_compensacao_jornada, characteristics_error = (
            _parse_employee_characteristics(request.POST)
        )
        cargo_id = (request.POST.get("cargo_id") or "").strip()
        sector_id = (request.POST.get("sector_id") or "").strip()
        work_schedule_id = (request.POST.get("work_schedule_id") or "").strip()

        if not nome_completo or not cargo_id or not sector_id:
            messages.error(
                request,
                "Nome completo, cargo e setor sao obrigatorios.",
            )
            return redirect("employee_edit_page", employee_id=employee.id)

        if characteristics_error:
            messages.error(request, characteristics_error)
            return redirect("employee_edit_page", employee_id=employee.id)

        if tipo == Employee.TYPE_DIRETO and not matricula:
            messages.error(
                request,
                "Matricula e obrigatoria para empregado direto.",
            )
            return redirect("employee_edit_page", employee_id=employee.id)

        duplicate = Employee.objects.filter(matricula=matricula).exclude(id=employee.id) if matricula else Employee.objects.none()
        if duplicate.exists():
            messages.error(request, "Ja existe empregado com essa matricula.")
            return redirect("employee_edit_page", employee_id=employee.id)

        cargo = active_cargos.filter(id=cargo_id).first()
        if (
            not cargo
            and current_cargo
            and current_cargo.deactivated_at
            and str(current_cargo.id) == cargo_id
        ):
            cargo = current_cargo

        if not cargo:
            messages.error(request, "Selecione um cargo valido.")
            return redirect("employee_edit_page", employee_id=employee.id)

        sector = active_sectors.filter(id=sector_id).first()
        if (
            not sector
            and current_sector
            and current_sector.deactivated_at
            and str(current_sector.id) == sector_id
        ):
            sector = current_sector

        if not sector:
            messages.error(request, "Selecione um setor valido.")
            return redirect("employee_edit_page", employee_id=employee.id)

        work_schedule, schedule_error = _get_work_schedule_or_error(work_schedule_id)
        if schedule_error:
            messages.error(request, schedule_error)
            return redirect("employee_edit_page", employee_id=employee.id)

        previous_tipo = employee.tipo
        previous_cargo = employee.cargo
        previous_sector = employee.sector
        previous_work_schedule = employee.work_schedule
        previous_cargo_id = previous_cargo.id if previous_cargo else None
        previous_sector_id = previous_sector.id if previous_sector else None
        previous_work_schedule_id = (
            previous_work_schedule.id if previous_work_schedule else None
        )
        next_cargo_id = cargo.id if cargo else None
        next_sector_id = sector.id if sector else None
        next_work_schedule_id = work_schedule.id if work_schedule else None

        cargo_changed = previous_cargo_id != next_cargo_id
        sector_changed = previous_sector_id != next_sector_id
        work_schedule_changed = previous_work_schedule_id != next_work_schedule_id
        tipo_changed = previous_tipo != tipo
        has_allocation_change = (
            cargo_changed or sector_changed or work_schedule_changed or tipo_changed
        )

        effective_date = None
        if has_allocation_change:
            raw_effective_date = (request.POST.get("change_effective_date") or "").strip()
            if not raw_effective_date:
                messages.error(
                    request,
                    "Informe a data de vigencia para alteracao de setor/escala.",
                )
                return redirect("employee_edit_page", employee_id=employee.id)

            try:
                effective_date = date.fromisoformat(raw_effective_date)
            except ValueError:
                messages.error(request, "Informe uma data de vigencia valida.")
                return redirect("employee_edit_page", employee_id=employee.id)

            minimum_effective_date = timezone.localdate() - timedelta(days=30)
            if effective_date < minimum_effective_date:
                messages.error(
                    request,
                    "Nao e permitido registrar alteracoes de cargo/setor/escala com vigencia superior a 30 dias no passado.",
                )
                return redirect("employee_edit_page", employee_id=employee.id)

        if has_allocation_change and effective_date:
            closed_month = _find_first_closed_competence_month(effective_date)
            if closed_month:
                messages.error(
                    request,
                    "Nao e permitido registrar ajuste retroativo com vigencia em competencia encerrada "
                    f"({_format_competence_month_label(closed_month)}).",
                )
                return redirect("employee_edit_page", employee_id=employee.id)
        employee.matricula = matricula or None
        employee.nome_completo = nome_completo
        employee.tipo = tipo
        employee.cargo = cargo
        employee.sector = sector
        employee.work_schedule = work_schedule
        # Regime de compensacao nao e atualizado nesta tela.
        # A mudanca deve ocorrer via registro de evento.
        employee.save(
            update_fields=[
                "matricula",
                "nome_completo",
                "tipo",
                "cargo",
                "sector",
                "work_schedule",
            ]
        )

        if has_allocation_change and effective_date:
            change_notes_parts = []
            if cargo_changed:
                previous_cargo_label = previous_cargo.nome if previous_cargo else "Sem cargo"
                next_cargo_label = cargo.nome if cargo else "Sem cargo"
                change_notes_parts.append(
                    f"Cargo: {previous_cargo_label} -> {next_cargo_label}"
                )
            if tipo_changed:
                change_notes_parts.append(f"Tipo: {previous_tipo} -> {tipo}")

            EmployeeEvent.objects.create(
                employee=employee,
                event_type=EmployeeEvent.EVENT_TYPE_ALLOCATION_CHANGE,
                effective_date=effective_date,
                notes=" | ".join(change_notes_parts),
                previous_sector=previous_sector,
                new_sector=sector,
                previous_work_schedule=previous_work_schedule,
                new_work_schedule=work_schedule,
            )

        messages.success(request, "Empregado atualizado com sucesso.")
        return redirect("employee_edit_page", employee_id=employee.id)

    return render(
        request,
        "employee_edit.html",
        {
            "employee": employee,
            "active_cargos": active_cargos,
            "active_sectors": active_sectors,
            "work_schedules": work_schedules,
            "sector_history": sector_history,
            "work_schedule_history": work_schedule_history,
            "employee_type_choices": Employee.TYPE_CHOICES,
            "compensation_regime_choices": Employee.REGIME_COMPENSACAO_JORNADA_CHOICES,
            "calendar_exception_events": calendar_exception_events,
            "selected_tab": selected_tab,
            "active_areas": active_areas,
            "active_departments": active_departments,
        },
    )


@require_POST
def employee_deactivate(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id, deactivated_at__isnull=True)
    employee.deactivated_at = timezone.now()
    employee.save(update_fields=["deactivated_at"])
    messages.success(request, "Empregado desativado com sucesso.")
    return redirect("employees_page")


@require_POST
def employee_activate(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id, deactivated_at__isnull=False)
    employee.deactivated_at = None
    employee.save(update_fields=["deactivated_at"])
    messages.success(request, "Empregado ativado com sucesso.")
    return redirect("employees_page")


def cargos_page(request):
    if request.method == "POST":
        nome = (request.POST.get("nome") or "").strip()

        if not nome:
            messages.error(request, "Nome do cargo e obrigatorio.")
            return redirect("cargos_page")

        if Cargo.objects.filter(nome__iexact=nome, deactivated_at__isnull=True).exists():
            messages.error(request, "Ja existe um cargo com esse nome.")
            return redirect("cargos_page")

        Cargo.objects.create(nome=nome)
        messages.success(request, "Cargo cadastrado com sucesso.")
        return redirect("cargos_page")

    cargos = (
        Cargo.objects.annotate(employee_count=Count("employees"))
        .order_by("deactivated_at", "nome")
    )
    active_count = cargos.filter(deactivated_at__isnull=True).count()
    deactivated_count = cargos.filter(deactivated_at__isnull=False).count()
    return render(
        request,
        "cargos.html",
        {
            "cargos": cargos,
            "active_count": active_count,
            "deactivated_count": deactivated_count,
        },
    )


def cargo_edit_page(request, cargo_id):
    cargo = get_object_or_404(Cargo, id=cargo_id, deactivated_at__isnull=True)

    if request.method == "POST":
        nome = (request.POST.get("nome") or "").strip()

        if not nome:
            messages.error(request, "Nome do cargo e obrigatorio.")
            return redirect("cargo_edit_page", cargo_id=cargo.id)

        duplicate = Cargo.objects.filter(
            nome__iexact=nome,
            deactivated_at__isnull=True,
        ).exclude(id=cargo.id)
        if duplicate.exists():
            messages.error(request, "Ja existe um cargo com esse nome.")
            return redirect("cargo_edit_page", cargo_id=cargo.id)

        cargo.nome = nome
        cargo.save(update_fields=["nome"])
        messages.success(request, "Cargo atualizado com sucesso.")
        return redirect("cargos_page")

    return render(request, "cargo_edit.html", {"cargo": cargo})


@require_POST
def cargo_deactivate(request, cargo_id):
    cargo = get_object_or_404(Cargo, id=cargo_id, deactivated_at__isnull=True)
    cargo.deactivated_at = timezone.now()
    cargo.save(update_fields=["deactivated_at"])
    messages.success(request, "Cargo desativado com sucesso.")
    return redirect("cargos_page")


@require_POST
def cargo_activate(request, cargo_id):
    cargo = get_object_or_404(Cargo, id=cargo_id, deactivated_at__isnull=False)
    cargo.deactivated_at = None
    cargo.save(update_fields=["deactivated_at"])
    messages.success(request, "Cargo ativado com sucesso.")
    return redirect("cargos_page")


def sectors_page(request):
    if request.method == "POST":
        create_without_department = (
            (request.POST.get("create_without_department") or "").strip() == "1"
        )
        if create_without_department:
            nome = (request.POST.get("nome") or "").strip()
            if not nome:
                messages.error(request, "Nome do setor e obrigatorio.")
                return redirect("sectors_page")
            if Sector.objects.filter(
                department__isnull=True,
                nome__iexact=nome,
                deactivated_at__isnull=True,
            ).exists():
                messages.error(request, "Ja existe um setor sem departamento com esse nome.")
                return redirect("sectors_page")
            Sector.objects.create(nome=nome, department=None)
            messages.success(request, "Setor sem departamento cadastrado com sucesso.")
            return redirect("sectors_page")
        else:
            _, department, nome, taxonomy_error = _parse_sector_taxonomy_or_error(request.POST)
            if taxonomy_error:
                messages.error(request, taxonomy_error)
                return redirect("sectors_page")

            if Sector.objects.filter(
                department=department,
                nome__iexact=nome,
                deactivated_at__isnull=True,
            ).exists():
                messages.error(request, "Ja existe um setor com esse nome neste departamento.")
                return redirect("sectors_page")

            Sector.objects.create(nome=nome, department=department)
            messages.success(request, "Setor cadastrado com sucesso.")
            return redirect("sectors_page")

    sectors = (
        Sector.objects.select_related("department__area")
        .all()
        .order_by("deactivated_at", "department__area__nome", "department__nome", "nome")
    )
    all_areas = Area.objects.all().order_by("deactivated_at", "nome")
    all_departments = Department.objects.select_related("area").all().order_by(
        "deactivated_at",
        "area__nome",
        "nome",
    )
    department_ids_with_sector = set(
        sectors.filter(department__isnull=False).values_list("department_id", flat=True)
    )
    area_ids_with_sector = set(
        sectors.filter(department__area__isnull=False).values_list(
            "department__area_id",
            flat=True,
        )
    )
    departments_without_sectors = all_departments.exclude(id__in=department_ids_with_sector)
    areas_without_sectors = all_areas.exclude(id__in=area_ids_with_sector)
    active_count = sectors.filter(deactivated_at__isnull=True).count()
    deactivated_count = sectors.filter(deactivated_at__isnull=False).count()
    active_areas, active_departments = _active_taxonomy()
    return render(
        request,
        "sectors.html",
        {
            "sectors": sectors,
            "active_count": active_count,
            "deactivated_count": deactivated_count,
            "active_areas": active_areas,
            "active_departments": active_departments,
            "departments_without_sectors": departments_without_sectors,
            "areas_without_sectors": areas_without_sectors,
        },
    )


def work_schedules_page(request):
    if request.method == "POST":
        form_type = (request.POST.get("form_type") or "work_schedule").strip()

        if form_type == "calendar":
            calendar_name = (request.POST.get("calendar_name") or "").strip()
            if not calendar_name:
                messages.error(request, "Nome do calendario e obrigatorio.")
                return redirect("work_schedules_page")

            if WorkCalendar.objects.filter(nome__iexact=calendar_name).exists():
                messages.error(request, "Ja existe um calendario com esse nome.")
                return redirect("work_schedules_page")

            WorkCalendar.objects.create(nome=calendar_name)
            messages.success(request, "Calendario cadastrado com sucesso.")
            return redirect("work_schedules_page")

        if form_type == "special_period":
            calendar_id = (request.POST.get("calendar_id") or "").strip()
            calendar, calendar_error = _get_work_calendar_or_error(calendar_id)
            if calendar_error:
                messages.error(request, calendar_error)
                return redirect("work_schedules_page")

            latest_closure = _get_latest_closure_for_calendar(calendar.id)
            if latest_closure:
                month_label = _format_competence_month_label(latest_closure.competence_month)
                messages.error(
                    request,
                    f"Nao e permitido alterar este calendario, pois ele ja foi usado em mes fechado ({month_label}). Crie um novo calendario.",
                )
                return redirect("work_schedules_page")

            period_type = (request.POST.get("period_type") or "").strip()
            valid_types = {choice for choice, _ in WorkCalendarPeriod.TYPE_CHOICES}
            if period_type not in valid_types:
                messages.error(request, "Selecione um tipo valido de periodo.")
                return redirect("work_schedules_page")

            start_date, end_date, date_error = _parse_special_period_dates(request.POST)
            if date_error:
                messages.error(request, date_error)
                return redirect("work_schedules_page")

            today = timezone.localdate()
            if start_date < today or end_date < today:
                messages.error(
                    request,
                    "Nao e permitido incluir excecoes com datas que ja passaram.",
                )
                return redirect("work_schedules_page")

            description = (request.POST.get("description") or "").strip()

            WorkCalendarPeriod.objects.create(
                calendar=calendar,
                period_type=period_type,
                start_date=start_date,
                end_date=end_date,
                description=description,
            )
            messages.success(request, "Periodo especial cadastrado com sucesso.")
            return redirect("work_schedules_page")

        if form_type == "work_schedule_edit":
            work_schedule_id = (request.POST.get("work_schedule_id") or "").strip()
            work_schedule = WorkSchedule.objects.filter(id=work_schedule_id).first()
            if not work_schedule:
                messages.error(request, "Selecione uma escala valida.")
                return redirect("work_schedules_page")

            latest_closure = _get_latest_closure_for_schedule(work_schedule.id)
            if latest_closure:
                month_label = _format_competence_month_label(latest_closure.competence_month)
                messages.error(
                    request,
                    f"Nao e permitido alterar esta escala, pois ela ja foi usada em mes fechado ({month_label}). Crie uma nova escala.",
                )
                return redirect("work_schedules_page")

            nome = (request.POST.get("nome") or "").strip()
            if not nome:
                messages.error(request, "Nome da escala e obrigatorio.")
                return redirect("work_schedules_page")

            duplicate = WorkSchedule.objects.filter(nome__iexact=nome).exclude(
                id=work_schedule.id
            )
            if duplicate.exists():
                messages.error(request, "Ja existe uma escala com esse nome.")
                return redirect("work_schedules_page")

            calendar_id = (request.POST.get("calendar_id") or "").strip()
            calendar, calendar_error = _get_work_calendar_or_error(calendar_id)
            if calendar_error:
                messages.error(request, calendar_error)
                return redirect("work_schedules_page")

            daily_hours, error = _parse_daily_hours(request.POST)
            if error:
                messages.error(request, error)
                return redirect("work_schedules_page")

            work_schedule.nome = nome
            work_schedule.calendar = calendar
            for field_name, _ in WEEKDAY_HOUR_FIELDS:
                setattr(work_schedule, field_name, daily_hours[field_name])
            work_schedule.save(
                update_fields=[
                    "nome",
                    "calendar",
                    *SCHEDULE_FIELDS_IN_WEEKDAY_ORDER,
                ]
            )
            messages.success(request, "Escala atualizada com sucesso.")
            return redirect("work_schedules_page")

        nome = (request.POST.get("nome") or "").strip()

        if not nome:
            messages.error(request, "Nome da escala e obrigatorio.")
            return redirect("work_schedules_page")

        if WorkSchedule.objects.filter(nome__iexact=nome).exists():
            messages.error(request, "Ja existe uma escala com esse nome.")
            return redirect("work_schedules_page")

        calendar_id = (request.POST.get("calendar_id") or "").strip()
        calendar, calendar_error = _get_work_calendar_or_error(calendar_id)
        if calendar_error:
            messages.error(request, calendar_error)
            return redirect("work_schedules_page")

        daily_hours, error = _parse_daily_hours(request.POST)
        if error:
            messages.error(request, error)
            return redirect("work_schedules_page")

        WorkSchedule.objects.create(nome=nome, calendar=calendar, **daily_hours)
        messages.success(request, "Escala cadastrada com sucesso.")
        return redirect("work_schedules_page")

    work_schedules = WorkSchedule.objects.select_related("calendar").order_by("nome")
    work_calendars = WorkCalendar.objects.all().order_by("nome")

    selected_calendar_id = (request.GET.get("calendar_id") or "").strip()
    selected_calendar = None
    if selected_calendar_id:
        selected_calendar = work_calendars.filter(id=selected_calendar_id).first()

    all_special_periods = WorkCalendarPeriod.objects.select_related("calendar").order_by(
        "start_date",
        "end_date",
        "id",
    )
    if selected_calendar:
        all_special_periods = all_special_periods.filter(calendar=selected_calendar)
    else:
        all_special_periods = all_special_periods.none()

    today = timezone.localdate()
    old_period_cutoff = today - timedelta(days=60)
    recent_special_periods = all_special_periods.filter(end_date__gte=old_period_cutoff)
    old_special_periods = all_special_periods.filter(end_date__lt=old_period_cutoff)

    return render(
        request,
        "work_schedules.html",
        {
            "work_schedules": work_schedules,
            "work_calendars": work_calendars,
            "recent_special_periods": recent_special_periods,
            "old_special_periods": old_special_periods,
            "selected_calendar": selected_calendar,
            "selected_calendar_id": str(selected_calendar.id) if selected_calendar else "",
            "today": today,
            "today_iso": today.isoformat(),
            "old_period_cutoff": old_period_cutoff,
        },
    )


@require_POST
def area_create(request):
    nome = (request.POST.get("nome") or "").strip()
    if not nome:
        messages.error(request, "Nome da area e obrigatorio.")
        return redirect("sectors_page")

    if Area.objects.filter(nome__iexact=nome, deactivated_at__isnull=True).exists():
        messages.error(request, "Ja existe uma area com esse nome.")
        return redirect("sectors_page")

    Area.objects.create(nome=nome)
    messages.success(request, "Area cadastrada com sucesso.")
    return redirect("sectors_page")


@require_POST
def area_edit(request, area_id):
    area = get_object_or_404(Area, id=area_id)
    nome = (request.POST.get("nome") or "").strip()
    if not nome:
        messages.error(request, "Nome da area e obrigatorio.")
        return redirect("sectors_page")

    duplicate = Area.objects.filter(
        nome__iexact=nome,
        deactivated_at__isnull=True,
    ).exclude(id=area.id)
    if duplicate.exists():
        messages.error(request, "Ja existe uma area com esse nome.")
        return redirect("sectors_page")

    area.nome = nome
    area.save(update_fields=["nome"])
    messages.success(request, "Area atualizada com sucesso.")
    return redirect("sectors_page")


@require_POST
def area_deactivate(request, area_id):
    area = get_object_or_404(Area, id=area_id, deactivated_at__isnull=True)
    now = timezone.now()
    with transaction.atomic():
        area.deactivated_at = now
        area.save(update_fields=["deactivated_at"])
        Department.objects.filter(area=area, deactivated_at__isnull=True).update(
            deactivated_at=now
        )
        Sector.objects.filter(
            department__area=area,
            deactivated_at__isnull=True,
        ).update(deactivated_at=now)
    messages.success(request, "Area desativada com sucesso.")
    return redirect("sectors_page")


@require_POST
def area_activate(request, area_id):
    area = get_object_or_404(Area, id=area_id, deactivated_at__isnull=False)
    area.deactivated_at = None
    area.save(update_fields=["deactivated_at"])
    messages.success(request, "Area ativada com sucesso.")
    return redirect("sectors_page")


@require_POST
def department_create(request):
    area_id = (request.POST.get("area_id") or "").strip()
    nome = (request.POST.get("nome") or "").strip()

    if not area_id:
        messages.error(request, "Selecione uma area para o departamento.")
        return redirect("sectors_page")
    if not nome:
        messages.error(request, "Nome do departamento e obrigatorio.")
        return redirect("sectors_page")

    area = Area.objects.filter(id=area_id, deactivated_at__isnull=True).first()
    if not area:
        messages.error(request, "Selecione uma area valida.")
        return redirect("sectors_page")

    if Department.objects.filter(
        area=area,
        nome__iexact=nome,
        deactivated_at__isnull=True,
    ).exists():
        messages.error(request, "Ja existe um departamento com esse nome nesta area.")
        return redirect("sectors_page")

    Department.objects.create(area=area, nome=nome)
    messages.success(request, "Departamento cadastrado com sucesso.")
    return redirect("sectors_page")


@require_POST
def department_edit(request, department_id):
    department = get_object_or_404(Department.objects.select_related("area"), id=department_id)
    area_id = (request.POST.get("area_id") or "").strip()
    nome = (request.POST.get("nome") or "").strip()

    if not area_id:
        messages.error(request, "Selecione uma area para o departamento.")
        return redirect("sectors_page")
    if not nome:
        messages.error(request, "Nome do departamento e obrigatorio.")
        return redirect("sectors_page")

    area = Area.objects.filter(id=area_id, deactivated_at__isnull=True).first()
    if not area:
        messages.error(request, "Selecione uma area valida.")
        return redirect("sectors_page")

    duplicate = Department.objects.filter(
        area=area,
        nome__iexact=nome,
        deactivated_at__isnull=True,
    ).exclude(id=department.id)
    if duplicate.exists():
        messages.error(request, "Ja existe um departamento com esse nome nesta area.")
        return redirect("sectors_page")

    department.area = area
    department.nome = nome
    department.save(update_fields=["area", "nome"])
    messages.success(request, "Departamento atualizado com sucesso.")
    return redirect("sectors_page")


@require_POST
def department_deactivate(request, department_id):
    department = get_object_or_404(Department, id=department_id, deactivated_at__isnull=True)
    now = timezone.now()
    with transaction.atomic():
        department.deactivated_at = now
        department.save(update_fields=["deactivated_at"])
        Sector.objects.filter(
            department=department,
            deactivated_at__isnull=True,
        ).update(deactivated_at=now)
    messages.success(request, "Departamento desativado com sucesso.")
    return redirect("sectors_page")


@require_POST
def department_activate(request, department_id):
    department = get_object_or_404(Department, id=department_id, deactivated_at__isnull=False)
    if department.area.deactivated_at:
        messages.error(request, "Ative a area antes de ativar este departamento.")
        return redirect("sectors_page")
    department.deactivated_at = None
    department.save(update_fields=["deactivated_at"])
    messages.success(request, "Departamento ativado com sucesso.")
    return redirect("sectors_page")


def sector_edit_page(request, sector_id):
    sector = get_object_or_404(
        Sector.objects.select_related("department__area"),
        id=sector_id,
        deactivated_at__isnull=True,
    )

    if request.method == "POST":
        _, department, nome, taxonomy_error = _parse_sector_taxonomy_or_error(request.POST)
        if taxonomy_error:
            messages.error(request, taxonomy_error)
            return redirect("sector_edit_page", sector_id=sector.id)

        duplicate = Sector.objects.filter(
            department=department,
            nome__iexact=nome,
            deactivated_at__isnull=True,
        ).exclude(id=sector.id)
        if duplicate.exists():
            messages.error(request, "Ja existe um setor com esse nome neste departamento.")
            return redirect("sector_edit_page", sector_id=sector.id)

        sector.nome = nome
        sector.department = department
        sector.save(update_fields=["nome", "department"])
        messages.success(request, "Setor atualizado com sucesso.")
        return redirect("sectors_page")

    active_areas, active_departments = _active_taxonomy()
    return render(
        request,
        "sector_edit.html",
        {
            "sector": sector,
            "active_areas": active_areas,
            "active_departments": active_departments,
        },
    )


@require_POST
def sector_deactivate(request, sector_id):
    sector = get_object_or_404(Sector, id=sector_id, deactivated_at__isnull=True)
    sector.deactivated_at = timezone.now()
    sector.save(update_fields=["deactivated_at"])
    messages.success(request, "Setor desativado com sucesso.")
    return redirect("sectors_page")


@require_POST
def sector_activate(request, sector_id):
    sector = get_object_or_404(Sector, id=sector_id, deactivated_at__isnull=False)
    sector.deactivated_at = None
    sector.save(update_fields=["deactivated_at"])
    messages.success(request, "Setor ativado com sucesso.")
    return redirect("sectors_page")


def database_settings_page(request):
    runtime_config = _load_db_runtime_config()
    mode = str(
        runtime_config.get("mode")
        or os.getenv("DATABASE_MODE", "arquivo")
    ).strip().lower()
    if mode not in {"online", "arquivo"}:
        mode = "arquivo"

    default_db_path = str(Path(__file__).resolve().parent.parent / "data" / "rh.db")
    db_file_path = str(
        runtime_config.get("db_file_path")
        or os.getenv("DB_FILE_PATH", "")
        or default_db_path
    ).strip()
    import_errors = request.session.pop("db_settings_employee_import_errors", None)
    import_result_count = request.session.pop("db_settings_employee_import_result_count", None)

    if request.method == "POST":
        form_type = (request.POST.get("form_type") or "").strip()
        if form_type == "employee_import_csv_db_settings":
            imported_count, errors = _import_employees_from_csv(request.FILES.get("csv_file"))
            if errors:
                request.session["db_settings_employee_import_errors"] = errors
                request.session["db_settings_employee_import_result_count"] = imported_count
            elif imported_count:
                messages.success(
                    request,
                    f"Importacao concluida: {imported_count} empregado(s) importado(s).",
                )
            return redirect("database_settings_page")

        posted_mode = (request.POST.get("mode") or "").strip().lower()
        posted_db_file_path = (request.POST.get("db_file_path") or "").strip()
        wants_to_open_location = (request.POST.get("open_location") or "").strip() == "1"
        wants_to_create_new_db = (request.POST.get("create_new_db") or "").strip() == "1"

        if posted_mode not in {"online", "arquivo"}:
            messages.error(request, "Modo de banco invalido.")
            return redirect("database_settings_page")

        if wants_to_open_location:
            if posted_mode != "arquivo":
                messages.error(request, "A abertura de localizacao so funciona no modo arquivo.")
                return redirect("database_settings_page")
            if not posted_db_file_path:
                messages.error(request, "Informe o caminho do arquivo .db para abrir a localizacao.")
                return redirect("database_settings_page")
            if not posted_db_file_path.lower().endswith(".db"):
                messages.error(request, "O caminho informado deve apontar para um arquivo .db.")
                return redirect("database_settings_page")

            db_candidate = Path(posted_db_file_path).expanduser()
            folder_candidate = db_candidate.parent
            if not folder_candidate.exists() or not folder_candidate.is_dir():
                messages.error(request, "A pasta informada no caminho do banco nao existe.")
                return redirect("database_settings_page")

            try:
                _open_location_in_file_manager(posted_db_file_path)
            except OSError:
                messages.error(request, "Nao foi possivel abrir a localizacao informada.")
            else:
                messages.success(request, "Localizacao aberta no explorador de arquivos.")
            return redirect("database_settings_page")

        if wants_to_create_new_db:
            if posted_mode != "arquivo":
                messages.error(request, "A criacao de novo banco so funciona no modo arquivo.")
                return redirect("database_settings_page")
            if not posted_db_file_path:
                messages.error(request, "Informe o caminho do arquivo .db para criar o novo banco.")
                return redirect("database_settings_page")
            if not posted_db_file_path.lower().endswith(".db"):
                messages.error(request, "O caminho informado deve apontar para um arquivo .db.")
                return redirect("database_settings_page")

            db_candidate = Path(posted_db_file_path).expanduser()
            if db_candidate.exists():
                messages.error(
                    request,
                    "Ja existe um arquivo nesse caminho. Informe outro caminho para criar um banco novo.",
                )
                return redirect("database_settings_page")

            created, error_detail = _create_clean_sqlite_db(posted_db_file_path)
            if not created:
                messages.error(
                    request,
                    f"Nao foi possivel criar o novo banco. Detalhe: {error_detail}",
                )
                return redirect("database_settings_page")

            _save_db_runtime_config("arquivo", str(db_candidate))
            messages.success(
                request,
                "Novo banco criado com sucesso. Configuracao atualizada para usar esse arquivo apos reiniciar o aplicativo.",
            )
            return redirect("database_settings_page")

        if posted_mode == "arquivo":
            if not posted_db_file_path:
                messages.error(request, "Informe o caminho do arquivo .db.")
                return redirect("database_settings_page")
            if not posted_db_file_path.lower().endswith(".db"):
                messages.error(request, "O arquivo informado deve ter extensao .db.")
                return redirect("database_settings_page")
            db_file_path = posted_db_file_path
        else:
            db_file_path = default_db_path

        _save_db_runtime_config(posted_mode, db_file_path)
        messages.success(
            request,
            "Configuracao salva. Reinicie o aplicativo para aplicar o novo banco de dados.",
        )
        return redirect("database_settings_page")

    context = {
        "database_mode": mode,
        "db_file_path": _portable_db_path_for_storage(db_file_path),
        "employee_import_errors": import_errors or [],
        "employee_import_result_count": import_result_count,
    }
    return render(request, "database_settings.html", context)


def employees_csv_template_download(request):
    csv_lines = [
        "matricula,admissao,nome,setor,cargo,tipo",
        "12345,14/05/2026,Joao da Silva,Usinagem,Operador,direto",
        "99999,02/01/2026,Maria Souza,Logistica,Analista,indireto",
    ]
    content = "\n".join(csv_lines) + "\n"
    response = HttpResponse(content, content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_empregados.csv"'
    return response






















